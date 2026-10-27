# -*- coding: utf-8 -*-
"""SafePlum 무료 내려받기 양식(엑셀) 만들기 — python3 scripts/templates/make_xlsx.py
결과: assets/files/templates/*.xlsx
"""
import json, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from xl_kit import *  # noqa
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "assets/files/templates"
OUT.mkdir(parents=True, exist_ok=True)


# =====================================================================
# 1. 위험성평가표
# =====================================================================
def risk():
    wb = new_book()
    cols = ["No", "공정(작업)", "세부 작업", "위험 분류", "원인", "유해·위험요인 및 위험성", "관련 근거\n(법적 기준)", "현재 안전조치",
            "가능성\n(빈도)", "중대성\n(강도)", "위험성", "등급", "감소대책", "개선 후\n가능성", "개선 후\n중대성", "개선 후\n위험성", "개선 후\n등급",
            "담당자", "개선\n예정일", "완료일", "비고"]
    widths = [4.5, 11, 13, 9, 13, 28, 14, 18, 6.5, 6.5, 6.5, 8, 28, 6.5, 6.5, 6.5, 8, 9, 10, 10, 9]
    ws = sheet(wb, "위험성평가표", widths, landscape=True, tab=NAVY, one_page=True)
    L = len(cols)
    title(ws, "위험성평가표", 14, sub="「사업장 위험성평가에 관한 지침」(고용노동부 고시)에 따른 유해·위험요인 파악 → 위험성 결정 → 감소대책 → 재평가 기록")
    approval(ws, 1, 18, labels=("담당", "검토", "승인"), sign_h=30)
    # 기본 정보
    r = 4
    info = [("사업장", "A4:B4", "C4:F4"), ("평가 구분", "G4", "H4"), ("평가일", "I4:J4", "K4:M4"), ("평가자(참여자)", "N4:P4", "Q4:U4")]
    for lab, a, b in info:
        box(ws, a, lab, f(9.5, True, NAVY), CENTER, HEAD); box(ws, b, None, f(10), CENTER)
    dv_list(ws, "H4", ["최초평가", "정기평가", "수시평가", "상시평가"])
    ws["H4"].value = None
    ws.row_dimensions[4].height = 24
    legend(ws, 5, L, "흰 칸에 입력 · 회색 기울임 줄은 예시(지우고 쓰기) · 위험성·등급은 가능성×중대성으로 자동 계산 · 판단 기준은 '평가기준' 시트에서 사업장 기준으로 고쳐 쓰세요")
    header_row(ws, 6, cols, height=34)
    first, last = 7, 26
    grid(ws, first, last, 1, L, height=28)
    for rr in range(first, last + 1):
        ws.cell(row=rr, column=1, value=rr - first + 1)
        for cc in (6, 8, 13):
            ws.cell(row=rr, column=cc).alignment = LEFT
        ws.cell(row=rr, column=11, value=f'=IF(OR(I{rr}="",J{rr}=""),"",I{rr}*J{rr})')
        ws.cell(row=rr, column=12, value=f'=IF(K{rr}="","",INDEX(평가기준!$C$15:$C$19,MATCH(K{rr},평가기준!$A$15:$A$19,1)))')
        ws.cell(row=rr, column=16, value=f'=IF(OR(N{rr}="",O{rr}=""),"",N{rr}*O{rr})')
        ws.cell(row=rr, column=17, value=f'=IF(P{rr}="","",INDEX(평가기준!$C$15:$C$19,MATCH(P{rr},평가기준!$A$15:$A$19,1)))')
        for cc in (11, 12, 16, 17):
            ws.cell(row=rr, column=cc).fill = fill(SUB); ws.cell(row=rr, column=cc).font = f(9.5, True)
        for cc in (19, 20):
            ws.cell(row=rr, column=cc).number_format = "yyyy-mm-dd"
    example(ws, first, [None, "가공", "프레스 금형 교체", "기계적 요인", "금형 교체 중 슬라이드 하강", "금형 교체 중 슬라이드가 내려와 손이 끼일 위험",
                        "안전보건규칙 제104조", "작업 전 전원 차단(구두 지시)", 3, 4, None, None, "안전블록 설치·전원 잠금(LOTO) 표지 부착, 2인 1조 작업",
                        1, 4, None, None, "생산1팀", "2026-11-15", None, None])
    for cc, v in ((9, 3), (10, 4), (14, 1), (15, 4)):
        ws.cell(row=first, column=cc).value = v
    dv_list(ws, f"D{first}:D{last}", ["기계적 요인", "전기적 요인", "화학(물질)적 요인", "생물학적 요인", "작업특성 요인", "작업환경 요인", "기타"])
    dv_whole(ws, f"I{first}:I{last}", 1, 5, "가능성 1~5 (평가기준 시트)")
    dv_whole(ws, f"N{first}:N{last}", 1, 5, "개선 후 가능성 1~5")
    dv_whole(ws, f"J{first}:J{last}", 1, 4, "중대성 1~4 (평가기준 시트)")
    dv_whole(ws, f"O{first}:O{last}", 1, 4, "개선 후 중대성 1~4")
    dv_date(ws, f"S{first}:T{last}")
    for rng in (f"L{first}:L{last}", f"Q{first}:Q{last}"):
        color_text_rule(ws, rng, "매우 높음", "B42318")
        color_text_rule(ws, rng, "높음", RED)
        color_text_rule(ws, rng, "보통", ORANGE)
        color_text_rule(ws, rng, "매우 낮음", "A6E3C1", "05603A")
        color_text_rule(ws, rng, "낮음", "D1FADF", "05603A")
    # 허용 기준 초과인데 담당·예정일 비어 있으면 표시
    ws.conditional_formatting.add(f"R{first}:S{last}", FormulaRule(formula=[f'AND($K{first}<>"",$K{first}>평가기준!$C$22,R{first}="")'], fill=fill(YELLOW)))
    # 요약
    sr = last + 2
    box(ws, f"A{sr}:E{sr}", "허용 기준 초과(개선 전)", f(9.5, True, NAVY), CENTER, HEAD)
    box(ws, f"F{sr}", f'=COUNTIF(K{first}:K{last},">"&평가기준!$C$22)&"건"', f(10, True, RED), CENTER)
    box(ws, f"G{sr}:H{sr}", "개선 후 여전히 초과", f(9.5, True, NAVY), CENTER, HEAD)
    box(ws, f"I{sr}:K{sr}", f'=COUNTIF(P{first}:P{last},">"&평가기준!$C$22)&"건"', f(10, True, RED), CENTER)
    box(ws, f"L{sr}:M{sr}", "완료 / 대상", f(9.5, True, NAVY), CENTER, HEAD)
    box(ws, f"N{sr}:Q{sr}", f'=COUNT(T{first}:T{last})&" / "&COUNT(K{first}:K{last})', f(10, True), CENTER)
    ws.row_dimensions[sr].height = 24
    ws.print_title_rows = "6:6"
    ws.freeze_panes = "C7"
    ws.print_area = f"A1:U{sr}"

    # ---- 평가기준
    g = sheet(wb, "평가기준", [12, 16, 40, 10, 10, 10, 10], tab="F79009", one_page=True)
    title(g, "위험성 평가기준 (가능성 × 중대성)", 7, sub="예시 기준입니다. 사업장 위험성평가 실시 규정에 맞게 숫자와 설명을 고쳐 쓰면 '위험성평가표' 시트가 따라 계산됩니다.")
    header_row(g, 4, ["가능성", "구분", "판단 기준(예시)"])
    P = [(5, "매우 높음", "피해가 생길 가능성이 매우 높음 — 같은 작업에서 사고·아차사고가 자주 있었음"),
         (4, "높음", "피해가 생길 가능성이 높음 — 안전조치가 거의 없음"),
         (3, "보통", "부주의하면 피해가 생길 수 있음 — 안전조치가 일부 미흡"),
         (2, "낮음", "피해가 생길 가능성이 낮음 — 안전조치가 대부분 되어 있음"),
         (1, "매우 낮음", "피해가 생길 가능성이 거의 없음")]
    for i, (n, k, d) in enumerate(P):
        rr = 5 + i
        box(g, f"A{rr}", n, f(10, True), CENTER); box(g, f"B{rr}", k, f(10), CENTER); box(g, f"C{rr}", d, f(9.5), LEFT)
        g.row_dimensions[rr].height = 30
    header_row(g, 4, ["중대성", "구분", "판단 기준(예시)"], start_col=5)
    g.column_dimensions["E"].width = 10; g.column_dimensions["F"].width = 12; g.column_dimensions["G"].width = 40
    C = [(4, "최대", "사망 또는 영구 장해"), (3, "대", "휴업이 필요한 부상·질병"), (2, "중", "휴업 없이 치료가 필요한 부상·질병"), (1, "소", "응급처치 수준의 경미한 부상")]
    for i, (n, k, d) in enumerate(C):
        rr = 5 + i
        box(g, f"E{rr}", n, f(10, True), CENTER); box(g, f"F{rr}", k, f(10), CENTER); box(g, f"G{rr}", d, f(9.5), LEFT)
    # 등급표(최소 점수 오름차순 — MATCH 1 사용)
    g.cell(row=12, column=1, value="위험성 등급 (위험성 = 가능성 × 중대성)").font = f(11, True, NAVY)
    header_row(g, 14, ["최소 점수", "최대 점수", "등급", "조치 기준"], height=24)
    g.column_dimensions["D"].width = 30
    G = [(1, 3, "매우 낮음", "현재 상태 유지"), (4, 6, "낮음", "현재 상태 유지 · 필요 시 개선"), (8, 10, "보통", "계획을 세워 개선"),
         (12, 15, "높음", "빠른 시일 안에 개선"), (16, 20, "매우 높음", "즉시 개선 · 개선 전 작업 중지 검토")]
    gc = {"매우 낮음": ("A6E3C1", "05603A"), "낮음": ("D1FADF", "05603A"), "보통": (ORANGE, "FFFFFF"), "높음": (RED, "FFFFFF"), "매우 높음": ("B42318", "FFFFFF")}
    for i, (a, b, k, d) in enumerate(G):
        rr = 15 + i
        box(g, f"A{rr}", a, f(10, True), CENTER); box(g, f"B{rr}", b, f(10), CENTER)
        box(g, f"C{rr}", k, f(10, True, gc[k][1]), CENTER, gc[k][0]); box(g, f"D{rr}", d, f(9.5), LEFT)
        g.row_dimensions[rr].height = 22
    box(g, "A22:B22", "허용 기준(이 점수 초과 시 감소대책)", f(9.5, True, NAVY), CENTER, HEAD)
    box(g, "C22", 6, f(11, True, RED), CENTER, YELLOW)
    g.row_dimensions[22].height = 30
    # 매트릭스
    g.cell(row=24, column=1, value="위험성 매트릭스").font = f(11, True, NAVY)
    box(g, "A25", "가능성 \\ 중대성", f(9, True, NAVY), CENTER, HEAD)
    for j in range(4):
        box(g, f"{'BCDE'[j]}25", 4 - j, f(10, True, NAVY), CENTER, HEAD)
    for i in range(5):
        rr = 26 + i; p = 5 - i
        box(g, f"A{rr}", p, f(10, True, NAVY), CENTER, HEAD)
        for j in range(4):
            col = "BCDE"[j]
            box(g, f"{col}{rr}", f"=$A{rr}*{col}$25", f(10, True), CENTER)
        g.row_dimensions[rr].height = 22
    rng = "B26:E30"
    g.conditional_formatting.add(rng, FormulaRule(formula=["B26>=16"], fill=fill("B42318"), font=Font(name=FONT, color="FFFFFF", bold=True)))
    g.conditional_formatting.add(rng, FormulaRule(formula=["B26>=12"], fill=fill(RED), font=Font(name=FONT, color="FFFFFF", bold=True)))
    g.conditional_formatting.add(rng, FormulaRule(formula=["B26>=8"], fill=fill(ORANGE), font=Font(name=FONT, color="FFFFFF", bold=True)))
    g.conditional_formatting.add(rng, FormulaRule(formula=["B26>=4"], fill=fill("D1FADF")))
    g.conditional_formatting.add(rng, FormulaRule(formula=["B26>=1"], fill=fill("A6E3C1")))
    g.print_area = "A1:G31"

    guide_sheet(wb, "위험성평가표", [
        ("이렇게 쓰세요", ["1) '위험성평가표' 시트 맨 위에 사업장·평가 구분(최초·정기·수시·상시)·평가일·참여자를 적습니다.",
                       "2) 공정(작업)마다 유해·위험요인을 '~하다가 ~로 ~할 위험' 형태로 적고, 지금 하고 있는 안전조치를 적습니다.",
                       "3) 가능성(1~5)·중대성(1~4)을 넣으면 위험성(곱)과 등급이 자동으로 나옵니다.",
                       "4) 허용 기준(평가기준 시트 C22, 기본 6점)을 넘는 줄은 감소대책·담당자·개선 예정일을 적습니다. 비어 있으면 노란색으로 표시됩니다.",
                       "5) 개선을 마치면 개선 후 가능성·중대성을 다시 넣어 재평가하고 완료일을 적습니다."]),
        ("감소대책 우선순위", ["제거·대체 → 공학적 대책(방호장치·차단·환기) → 관리적 대책(절차·허가·교육·감시인) → 보호구 순으로 검토합니다.",
                          "'안전교육 실시', '보호구 착용'만 적기보다 무엇을 어떻게 바꾸는지 구체적으로 적어야 현장에서 통합니다."]),
        ("참여와 공유", ["위험성평가는 해당 작업 근로자가 참여해야 하고(법 제36조제2항), 결과는 작업 전 TBM 등으로 근로자에게 알립니다.",
                      "결과와 조치 사항은 기록해 3년간 보존합니다(시행규칙 제37조)."]),
    ], source="근거: 산업안전보건법 제36조, 시행규칙 제37조, 「사업장 위험성평가에 관한 지침」 — 평가기준·등급은 예시이며 사업장 규정에 맞게 고쳐 쓰세요. SafePlum에서 화면으로 작성하려면 safeplum.com/tools/risk/")
    wb.move_sheet("작성 안내", offset=-2)
    wb.active = 1
    p = OUT / "SafePlum_risk_assessment.xlsx"; wb.save(p); return p


# =====================================================================
# 2. TBM 일지
# =====================================================================
PRE_ITEMS = ["작업자 건강상태(음주·피로·질병) 확인", "작업에 맞는 보호구 착용 확인", "사용 장비·공구 작업 전 점검", "작업구역 정리·통로 확보",
             "위험요인·대책 복창(전원 숙지)", "비상연락·대피 경로 공유"]


def tbm():
    wb = new_book()
    ws = sheet(wb, "TBM 일지", [8, 13, 13, 12, 13, 9, 13, 13, 13], tab=NAVY, one_page=True)
    title(ws, "TBM(작업 전 안전점검 회의) 일지", 6, sub="작업 시작 전 10분 — 오늘 작업의 위험요인과 대책을 함께 확인하고 기록합니다")
    approval(ws, 1, 7, labels=("관리감독자", "안전관리자"), sign_h=30)
    rows = [("일시", "B4:D4", "E4", "작업 장소", "F4:I4"), ("작업명", "B5:D5", "E5", "작업 인원", "F5:I5"), ("진행자", "B6:D6", "E6", "협력업체", "F6:I6")]
    for i, (a, av, b, bl, bv) in enumerate(rows):
        rr = 4 + i
        box(ws, f"A{rr}", a, f(9.5, True, NAVY), CENTER, HEAD); box(ws, av, None, f(10), LEFT)
        box(ws, b, bl, f(9.5, True, NAVY), CENTER, HEAD); box(ws, bv, None, f(10), LEFT)
        ws.row_dimensions[rr].height = 24
    ws["B4"].number_format = "yyyy-mm-dd hh:mm"
    example(ws, 4, [None, "2026-10-27 08:00"]); example(ws, 5, [None, "2층 배관 용접 보수"]); example(ws, 6, [None, "반장 ○○○"])
    # 작업 내용 / 달라진 점
    box(ws, "A8:I8", "1. 오늘 작업 내용과 어제와 달라진 점", f(10.5, True, "FFFFFF"), LEFT, NAVY)
    box(ws, "A9", "작업\n내용", f(9.5, True, NAVY), CENTER, HEAD); box(ws, "B9:I9", None, f(10), LEFT_TOP)
    box(ws, "A10", "달라진\n점", f(9.5, True, NAVY), CENTER, HEAD); box(ws, "B10:I10", None, f(10), LEFT_TOP)
    ws.row_dimensions[9].height = 44; ws.row_dimensions[10].height = 36
    example(ws, 9, [None, "2층 기계실 배관 용접 보수, 주변 단열재 있음 — 화기작업 허가 후 작업"]); example(ws, 10, [None, "어제와 달리 오후 협력업체 도장 작업과 같은 구역 동시 작업"])
    # 위험요인
    box(ws, "A12:I12", "2. 위험요인과 대책 (작업 단계별)", f(10.5, True, "FFFFFF"), LEFT, NAVY)
    header_row(ws, 13, ["No", "작업 단계", "", "유해·위험요인", "", "위험도", "감소대책", "", ""], height=24)
    ws.merge_cells("B13:C13"); ws.merge_cells("D13:E13"); ws.merge_cells("G13:I13")
    for i in range(6):
        rr = 14 + i
        grid(ws, rr, rr, 1, 9, height=28)
        ws.merge_cells(f"B{rr}:C{rr}"); ws.merge_cells(f"D{rr}:E{rr}"); ws.merge_cells(f"G{rr}:I{rr}")
        ws.cell(row=rr, column=1, value=i + 1)
        for cc in (2, 4, 7):
            ws.cell(row=rr, column=cc).alignment = LEFT
    example(ws, 14, [None, "용접 준비", None, "용접 불티가 단열재에 튀어 화재", None, "상", "불티 방지포 설치, 소화기 2대 비치, 화재감시자 배치"])
    dv_list(ws, "F14:F19", ["상", "중", "하"])
    color_text_rule(ws, "F14:F19", "상", RED)
    color_text_rule(ws, "F14:F19", "중", ORANGE)
    # 작업 전 확인
    box(ws, "A21:I21", "3. 작업 전 확인 (확인하면 ✔ 또는 O)", f(10.5, True, "FFFFFF"), LEFT, NAVY)
    for i, it in enumerate(PRE_ITEMS):
        rr = 22 + i // 2; c0 = 1 if i % 2 == 0 else 6
        if c0 == 1:
            box(ws, f"A{rr}", None, f(11, True, GREEN), CENTER); box(ws, f"B{rr}:E{rr}", it, f(9.5), LEFT)
        else:
            box(ws, f"F{rr}", None, f(11, True, GREEN), CENTER); box(ws, f"G{rr}:I{rr}", it, f(9.5), LEFT)
        ws.row_dimensions[rr].height = 22
    dv_list(ws, "A22:A24", ["✔", "O", "X"]); dv_list(ws, "F22:F24", ["✔", "O", "X"])
    # 근로자 의견
    box(ws, "A26:I26", "4. 근로자 의견·건의 (작업 전 확인과 구분해 적기)", f(10.5, True, "FFFFFF"), LEFT, NAVY)
    box(ws, "A27:I27", None, f(10), LEFT_TOP); ws.row_dimensions[27].height = 40
    example(ws, 27, ["예: 2층 통로 조명이 어두워 이동 시 미끄러질 위험 — 조명 추가 요청 → 시설팀 전달"])
    # 서명
    box(ws, "A29:I29", "5. 참석자 확인 (위 내용을 듣고 이해했습니다)", f(10.5, True, "FFFFFF"), LEFT, NAVY)
    header_row(ws, 30, ["No", "소속", "성명", "서명", "", "No", "소속", "성명", "서명"], height=22)
    ws.merge_cells("D30:E30")
    for i in range(7):
        rr = 31 + i
        grid(ws, rr, rr, 1, 9, height=24)
        ws.merge_cells(f"D{rr}:E{rr}")
        ws.cell(row=rr, column=1, value=i + 1); ws.cell(row=rr, column=6, value=i + 8)
    legend(ws, 39, 9)
    ws.print_area = "A1:I39"

    guide_sheet(wb, "TBM 일지", [
        ("진행 순서 (10분)", ["① 인원·건강상태 확인 → ② 오늘 작업과 달라진 점 설명 → ③ 단계별 위험요인·대책 공유 → ④ 작업 전 확인 6항목 → ⑤ 근로자 의견 → ⑥ 지적확인·복창 후 서명",
                          "같은 작업을 반복해도 '어제와 달라진 점'(동시 작업·날씨·인원·장비 변경)을 꼭 적으세요. 사고는 달라진 곳에서 납니다."]),
        ("위험성평가와 연결", ["위험성평가표의 해당 공정 유해·위험요인과 감소대책을 옮겨 적으면 위험성평가 결과를 근로자에게 알리는 기록이 됩니다(시행규칙 제37조 참고).",
                          "TBM에서 나온 새 위험요인은 수시 위험성평가로 반영하세요."]),
    ], source="TBM 자체는 법정 서식이 없는 실무 기록입니다. 위험성평가 결과 공유·근로자 참여 기록으로 활용할 수 있습니다. 화면으로 작성: safeplum.com/tools/tbm/")
    wb.move_sheet("작성 안내", offset=-1); wb.active = 1
    p = OUT / "SafePlum_tbm_log.xlsx"; wb.save(p); return p


# =====================================================================
# 3. 순회점검표 + 지적사항 관리대장
# =====================================================================
PATROL = [
    ("정리정돈·통로", ["통로·비상구에 적치물이 없고 폭이 확보되어 있다", "바닥에 기름·물·자재가 없어 미끄러지거나 걸려 넘어질 위험이 없다", "자재는 무너지지 않게 정해진 높이로 쌓여 있다"]),
    ("추락·개구부", ["개구부·단부에 안전난간 또는 덮개가 있다", "사다리는 미끄럼 방지·고정 상태가 양호하다", "2m 이상 고소작업 시 안전대를 걸고 작업한다"]),
    ("전기", ["분전반 문이 닫혀 있고 표지·잠금이 되어 있다", "누전차단기가 작동하고 이동 전선 피복 손상이 없다", "임시 배선이 통로를 가로지르지 않거나 보호되어 있다"]),
    ("화재·화기", ["소화기가 정해진 위치에 있고 압력이 정상이다", "화기작업은 허가서·불티 방지 조치·화재감시자가 있다", "인화성 물질이 화기 주변에 없다"]),
    ("기계·설비", ["회전·끼임 부위에 방호덮개가 있다", "정비·청소 시 전원 차단과 잠금·표지(LOTO)를 한다", "비상정지 장치가 작동한다"]),
    ("하역·운반", ["지게차 통로와 보행 통로가 구분되어 있다", "지게차 작업 전 점검을 하고 유자격자가 운전한다", "중량물은 인양 도구·줄걸이 상태가 양호하다"]),
    ("보호구·화학물질", ["작업에 맞는 보호구를 지급하고 착용한다", "화학물질 용기에 경고표지가 있고 MSDS가 게시되어 있다", "국소배기장치가 작동한다"]),
]
STATES = ["양호", "불량", "해당 없음", "미확인"]


def patrol():
    wb = new_book()
    ws = sheet(wb, "순회점검표", [5, 13, 44, 9, 15, 15], tab=NAVY, one_page=True)
    title(ws, "작업장 순회점검표", 3, sub="구역마다 항목을 점검하고 결과를 고르세요. 불량은 '지적사항 관리대장'에 옮겨 담당·기한을 정합니다.")
    approval(ws, 1, 4, labels=("점검자", "확인"), sign_h=30)
    rows = [("점검일", "구역·공정"), ("점검자", "동행자·협력업체")]
    for i, (a, b) in enumerate(rows):
        rr = 4 + i
        box(ws, f"A{rr}:B{rr}", a, f(9.5, True, NAVY), CENTER, HEAD)
        box(ws, f"C{rr}", None, f(10), LEFT)
        box(ws, f"D{rr}", b, f(9, True, NAVY), CENTER, HEAD)
        box(ws, f"E{rr}:F{rr}", None, f(10), LEFT)
        ws.row_dimensions[rr].height = 26
    ws["C4"].number_format = "yyyy-mm-dd"; dv_date(ws, "C4")
    header_row(ws, 7, ["No", "분류", "점검 항목", "결과", "지적 내용·위치", ""], height=26); ws.merge_cells("E7:F7")
    r = 8; n = 1; first = r
    for cat, items in PATROL:
        ws.merge_cells(start_row=r, start_column=2, end_row=r + len(items) - 1, end_column=2)
        for it in items:
            grid(ws, r, r, 1, 6, height=24); ws.merge_cells(f"E{r}:F{r}")
            ws.cell(row=r, column=1, value=n)
            ws.cell(row=r, column=3, value=it).alignment = LEFT
            ws.cell(row=r, column=5).alignment = LEFT
            r += 1; n += 1
        c = ws.cell(row=r - len(items), column=2, value=cat); c.font = f(9.5, True, NAVY); c.fill = fill(SUB); c.alignment = CENTER
    for k in range(2):
        grid(ws, r, r, 1, 6, height=24); ws.merge_cells(f"E{r}:F{r}"); ws.cell(row=r, column=1, value=n); ws.cell(row=r, column=2, value="추가 항목").font = f(9, False, MUTED)
        ws.cell(row=r, column=3).alignment = LEFT; r += 1; n += 1
    last = r - 1
    dv_list(ws, f"D{first}:D{last}", STATES)
    color_text_rule(ws, f"D{first}:D{last}", "불량", RED)
    color_text_rule(ws, f"D{first}:D{last}", "미확인", "FEF0C7", "B54708")
    color_text_rule(ws, f"D{first}:D{last}", "양호", "D1FADF", "05603A")
    example(ws, first, [None, None, None, "양호"])
    # 집계
    sr = last + 2
    for i, s in enumerate(STATES):
        box(ws, f"{'ABCD'[0]}{sr + i}:{'B'}{sr + i}", s, f(9.5, True, NAVY), CENTER, HEAD)
        box(ws, f"C{sr + i}", f'=COUNTIF($D${first}:$D${last},"{s}")&"건"', f(10, True, RED if s == "불량" else INK), LEFT)
    box(ws, f"D{sr}:F{sr + 3}", f'=IF(COUNTIF($D${first}:$D${last},"미확인")+COUNTBLANK($D${first}:$D${last})>0,"미확인·빈 항목 "&(COUNTIF($D${first}:$D${last},"미확인")+COUNTBLANK($D${first}:$D${last}))&"건 — 확인 필요","모든 항목 확인 완료")', f(10, True, NAVY), CENTER, SUB)
    legend(ws, sr + 5, 6)
    ws.print_area = f"A1:F{sr + 5}"
    ws.print_title_rows = "7:7"

    # 관리대장
    lc = ["No", "점검일", "구역·위치", "지적 내용", "위험도", "조치 부서·담당", "조치 기한", "조치 내용", "완료일", "완료 증빙\n(사진 번호 등)", "상태", "경과일"]
    m = sheet(wb, "지적사항 관리대장", [5, 11, 14, 32, 7, 14, 11, 32, 11, 14, 10, 7], landscape=True, tab=RED)
    title(m, "지적사항 조치 관리대장", 9, sub="상태는 자동: 완료일이 있으면 '완료', 기한이 지나면 '기한 경과', 3일 이내면 '임박'")
    approval(m, 1, 10, labels=("담당", "승인"), sign_h=28)
    legend(m, 4, len(lc))
    header_row(m, 5, lc, height=30)
    f0, f1 = 6, 39
    grid(m, f0, f1, 1, len(lc), height=25)
    for rr in range(f0, f1 + 1):
        m.cell(row=rr, column=1, value=rr - f0 + 1)
        for cc in (4, 8):
            m.cell(row=rr, column=cc).alignment = LEFT
        for cc in (2, 7, 9):
            m.cell(row=rr, column=cc).number_format = "yyyy-mm-dd"
        m.cell(row=rr, column=11, value=f'=IF(D{rr}="","",IF(I{rr}<>"","완료",IF(G{rr}="","기한 미정",IF(G{rr}<TODAY(),"기한 경과",IF(G{rr}-TODAY()<=3,"임박","진행 중")))))')
        m.cell(row=rr, column=12, value=f'=IF(OR(B{rr}="",D{rr}=""),"",IF(I{rr}<>"",I{rr},TODAY())-B{rr})')
        m.cell(row=rr, column=12).number_format = '0"일"'
        for cc in (11, 12):
            m.cell(row=rr, column=cc).fill = fill(SUB)
    example(m, f0, [None, "2026-10-27", "가공동 통로", "통로에 팔레트가 쌓여 지게차·보행자 동선이 겹침", "중", "물류팀 ○○○", "2026-10-30", "팔레트 지정 위치로 이동, 바닥 구획선 재도색", "2026-10-29", "사진 1·2"])
    dv_list(m, f"E{f0}:E{f1}", ["상", "중", "하"])
    dv_date(m, f"B{f0}:B{f1}"); dv_date(m, f"G{f0}:G{f1}"); dv_date(m, f"I{f0}:I{f1}")
    color_text_rule(m, f"K{f0}:K{f1}", "기한 경과", RED)
    color_text_rule(m, f"K{f0}:K{f1}", "임박", ORANGE)
    color_text_rule(m, f"K{f0}:K{f1}", "완료", "D1FADF", "05603A")
    color_text_rule(m, f"E{f0}:E{f1}", "상", RED)
    sr = f1 + 2
    for i, (lab, frm) in enumerate([("전체", f'=COUNTA(D{f0}:D{f1})'), ("완료", f'=COUNTIF(K{f0}:K{f1},"완료")'), ("진행 중·임박", f'=COUNTIF(K{f0}:K{f1},"진행 중")+COUNTIF(K{f0}:K{f1},"임박")'), ("기한 경과", f'=COUNTIF(K{f0}:K{f1},"기한 경과")')]):
        box(m, f"{'BDFH'[i]}{sr}", lab, f(9.5, True, NAVY), CENTER, HEAD)
        box(m, f"{'CEGI'[i]}{sr}", frm, f(11, True, RED if lab == "기한 경과" else INK), CENTER)
    m.row_dimensions[sr].height = 24
    m.print_title_rows = "5:5"; m.freeze_panes = "C6"; m.print_area = f"A1:L{sr}"

    guide_sheet(wb, "순회점검표·지적사항 관리대장", [
        ("순회점검표", ["점검 항목은 일반 작업장 공통 예시입니다. 업종·구역에 맞게 항목을 고치고 '추가 항목' 줄을 쓰세요.",
                    "결과는 양호·불량·해당 없음·미확인 중에서 고릅니다. 비어 있거나 미확인인 항목 수가 아래에 표시됩니다."]),
        ("지적사항 관리대장", ["불량 항목을 한 줄씩 옮겨 담당·조치 기한을 정합니다. 완료일을 넣으면 상태가 '완료'로 바뀝니다.",
                          "완료 증빙(사진)은 SafePlum 개선 전·후 사진대지(safeplum.com/tools/improve/)로 함께 보관하면 편합니다."]),
        ("도급 사업장이라면", ["도급인은 작업장 순회점검을 해야 합니다(산업안전보건법 제64조제1항제2호, 시행규칙 제80조: 건설업 등은 2일에 1회 이상, 그 밖의 사업은 1주일에 1회 이상)."]),
    ], source="점검 항목은 실무 예시이며 법정 서식이 아닙니다. 순회점검 주기는 시행규칙 제80조를 사업장 업종에 맞게 확인하세요. 화면으로 작성: safeplum.com/tools/patrol/")
    wb.move_sheet("작성 안내", offset=-2); wb.active = 1
    p = OUT / "SafePlum_patrol_checklist.xlsx"; wb.save(p); return p


# =====================================================================
# 4. 안전보건교육 일지·참석자 서명부
# =====================================================================
def edu():
    eh = json.load(open(ROOT / "data/edu_hours.json", encoding="utf-8"))
    wb = new_book()
    ws = sheet(wb, "교육일지", [12, 14, 14, 12, 14, 14], tab=NAVY, one_page=True)
    title(ws, "안전보건교육 일지", 3, sub="교육 과정·시간·내용·강사·참석 인원을 기록하고 참석자 서명부와 함께 보관합니다")
    approval(ws, 1, 4, labels=("담당", "승인"), sign_h=30)
    rows = [("교육 과정", "B4:C4", "교육 일자", "E4:F4"), ("교육 대상", "B5:C5", "교육 시간", "E5:F5"), ("교육 장소", "B6:C6", "법정 시간", "E6:F6"),
            ("교육 방법", "B7:C7", "강사", "E7:F7"), ("대상 인원", "B8:C8", "참석 인원", "E8:F8")]
    for i, (a, av, b, bv) in enumerate(rows):
        rr = 4 + i
        box(ws, f"A{rr}", a, f(9.5, True, NAVY), CENTER, HEAD); box(ws, av, None, f(10), LEFT)
        box(ws, f"D{rr}", b, f(9.5, True, NAVY), CENTER, HEAD); box(ws, bv, None, f(10), LEFT)
        ws.row_dimensions[rr].height = 24
    dv_list(ws, "B4", ["정기교육", "채용 시 교육", "작업내용 변경 시 교육", "특별교육", "관리감독자 정기교육", "건설업 기초안전·보건교육", "MSDS 교육", "기타"])
    dv_list(ws, "B7", ["집체교육", "현장교육", "인터넷 원격교육", "시청각교육"])
    ws["E4"].number_format = "yyyy-mm-dd"; dv_date(ws, "E4")
    example(ws, 4, [None, "정기교육"]); example(ws, 5, [None, "생산직 근로자", None, None, "2시간 (09:00~11:00)"])
    example(ws, 6, [None, "본관 2층 교육장", None, None, "매반기 12시간 이상 중 이번 2시간"])
    example(ws, 8, [None, "25명", None, None, "23명 (미참석 2명: 휴가)"])
    box(ws, "A10:F10", "교육 내용", f(10.5, True, "FFFFFF"), LEFT, NAVY)
    box(ws, "A11:F11", None, f(10), LEFT_TOP); ws.row_dimensions[11].height = 150
    example(ws, 11, ["예:\n1. 끼임 사고 사례와 원인(설비 가동 중 이물 제거)\n2. 전원 차단·잠금·표지(LOTO) 절차 실습\n3. 우리 공정 위험성평가 결과와 감소대책 공유\n4. 질의응답"])
    box(ws, "A13:F13", "교육 자료·사진", f(10.5, True, "FFFFFF"), LEFT, NAVY)
    box(ws, "A14:F14", None, f(10), LEFT_TOP); ws.row_dimensions[14].height = 120
    box(ws, "A16:B16", "확인 의견", f(9.5, True, NAVY), CENTER, HEAD); box(ws, "C16:F16", None, f(10), LEFT)
    ws.row_dimensions[16].height = 36
    legend(ws, 18, 6)
    ws.print_area = "A1:F18"

    s = sheet(wb, "참석자 서명부", [6, 16, 14, 18, 6, 16, 14, 18], tab=GREEN)
    title(s, "안전보건교육 참석자 서명부", 8, sub="교육 과정:                          교육 일자:                          교육 시간:")
    header_row(s, 4, ["No", "소속", "성명", "서명", "No", "소속", "성명", "서명"], height=24)
    for i in range(25):
        rr = 5 + i
        grid(s, rr, rr, 1, 8, height=27)
        s.cell(row=rr, column=1, value=i + 1); s.cell(row=rr, column=5, value=i + 26)
    s.print_title_rows = "4:4"; s.print_area = "A1:H29"

    t = sheet(wb, "법정 교육시간", [26, 46, 34], tab="F79009")
    title(t, "안전보건교육 교육과정별 교육시간 (요약)", 3, sub=eh["source"] + " · 확인일 " + eh["checked"])
    r = 4
    for sec in eh["sections"][:2]:
        box(t, f"A{r}:C{r}", sec["title"] + "  (" + sec.get("basis", "") + ")", f(10.5, True, "FFFFFF"), LEFT, NAVY); r += 1
        header_row(t, r, ["교육과정", "교육대상", "교육시간"], height=22); r += 1
        for row in sec["rows"]:
            box(t, f"A{r}", row["course"], f(9.5, True), LEFT)
            box(t, f"B{r}", row["target"], f(9.5), LEFT)
            box(t, f"C{r}", row["hours"], f(9.5), LEFT)
            t.row_dimensions[r].height = max(22, 15 * (1 + max(len(row["target"]) // 32, row["hours"].count("\n") + len(row["hours"]) // 30)))
            r += 1
        for nt in sec.get("notes", []):
            t.merge_cells(f"A{r}:C{r}"); c = t.cell(row=r, column=1, value="※ " + nt); c.font = f(8.5, False, MUTED); c.alignment = LEFT
            t.row_dimensions[r].height = 30 if len(nt) > 80 else 18; r += 1
        r += 1
    c = t.cell(row=r, column=1, value="전체 원문(직무교육 등 포함)은 국가법령정보센터 시행규칙 별표 4에서 확인하세요. SafePlum 법정교육 시간 도구: safeplum.com/tools/edu-hours/")
    c.font = f(8.5, False, MUTED, True); t.merge_cells(f"A{r}:C{r}")
    t.print_area = f"A1:C{r}"

    guide_sheet(wb, "안전보건교육 일지", [
        ("이렇게 쓰세요", ["교육 과정을 고르고, 교육 시간에는 이번 교육 시간을, 법정 시간에는 '법정 교육시간' 시트의 해당 기준을 적습니다.",
                       "정기교육은 반기·연간 합계로 채우면 되므로 회차별 일지를 모아 합계를 확인하세요.",
                       "참석자 서명부는 일지와 함께 보관합니다. 미참석자는 사유와 보충 교육 일정을 확인 의견에 적습니다."]),
    ], source="교육시간: " + eh["source"] + " (확인일 " + eh["checked"] + "). 화면으로 작성: safeplum.com/tools/edu-log/")
    wb.move_sheet("작성 안내", offset=-3); wb.active = 1
    p = OUT / "SafePlum_edu_log.xlsx"; wb.save(p); return p


# =====================================================================
# 5. 아차사고 보고서·관리대장
# =====================================================================
def near_miss():
    wb = new_book()
    ws = sheet(wb, "아차사고 보고서", [12, 14, 14, 12, 14, 14], tab=NAVY, one_page=True)
    title(ws, "아차사고 보고서", 3, sub="다친 사람이 없어도 같은 원인이 반복되면 사고가 됩니다 — 발견한 사람 누구나 작성")
    approval(ws, 1, 4, labels=("팀장", "안전관리자"), sign_h=30)
    rows = [("발생 일시", "B4:C4", "발생 장소", "E4:F4"), ("보고자", "B5:C5", "보고일", "E5:F5"), ("작업 내용", "B6:C6", "관련 설비", "E6:F6"), ("유형", "B7:C7", "잠재 피해", "E7:F7")]
    for i, (a, av, b, bv) in enumerate(rows):
        rr = 4 + i
        box(ws, f"A{rr}", a, f(9.5, True, NAVY), CENTER, HEAD); box(ws, av, None, f(10), LEFT)
        box(ws, f"D{rr}", b, f(9.5, True, NAVY), CENTER, HEAD); box(ws, bv, None, f(10), LEFT)
        ws.row_dimensions[rr].height = 24
    TYPES = ["떨어짐", "넘어짐", "끼임", "부딪힘", "물체에 맞음", "무너짐", "깔림·뒤집힘", "절단·베임·찔림", "감전", "화재·폭발", "화학물질 누출·접촉", "기타"]
    dv_list(ws, "B7", TYPES); dv_list(ws, "E7", ["사망 가능", "중상 가능", "경상 가능", "재산 피해"])
    ws["B4"].number_format = "yyyy-mm-dd hh:mm"; ws["E5"].number_format = "yyyy-mm-dd"
    example(ws, 4, [None, "2026-10-27 10:20", None, None, "자재창고 랙 3열"]); example(ws, 7, [None, "물체에 맞음", None, None, "중상 가능"])
    sec = [("무엇이 있었나 (경위)", 70, "예: 지게차로 랙 상단 팔레트를 내리던 중 박스 1개가 떨어졌으나 아래에 사람이 없어 다치지 않음"),
           ("왜 일어났나 (원인)", 0, None), ("어떻게 막나 (재발방지 대책)", 70, "예: 상단 적재 높이 1.5m 이하 제한, 하역 중 랙 아래 출입 통제선, 팔레트 랩핑 상태 작업 전 점검 항목에 추가"),
           ("조치 결과 확인", 46, None)]
    r = 9
    for name, h, ex in sec:
        box(ws, f"A{r}:F{r}", name, f(10.5, True, "FFFFFF"), LEFT, NAVY); r += 1
        if name.startswith("왜"):
            for k, lab, e in (("인적", "사람", "예: 하역 반경 출입 통제 미흡"), ("물적", "설비·물질", "예: 팔레트 랩핑 불량"), ("관리적", "절차·관리", "예: 적재 높이 기준 없음")):
                box(ws, f"A{r}", f"{k} 원인\n({lab})", f(9, True, NAVY), CENTER, SUB); box(ws, f"B{r}:F{r}", None, f(10), LEFT)
                example(ws, r, [None, e]); ws.row_dimensions[r].height = 30; r += 1
        else:
            box(ws, f"A{r}:F{r}", None, f(10), LEFT_TOP); ws.row_dimensions[r].height = h
            if ex:
                example(ws, r, [ex])
            r += 1
        r += 1
    rr = r - 2
    box(ws, f"A{r}", "조치 담당", f(9.5, True, NAVY), CENTER, HEAD); box(ws, f"B{r}:C{r}", None, f(10), LEFT)
    box(ws, f"D{r}", "조치 기한", f(9.5, True, NAVY), CENTER, HEAD); box(ws, f"E{r}:F{r}", None, f(10), LEFT)
    ws.row_dimensions[r].height = 24; r += 1
    box(ws, f"A{r}", "공유 방법", f(9.5, True, NAVY), CENTER, HEAD); box(ws, f"B{r}:C{r}", None, f(10), LEFT)
    box(ws, f"D{r}", "위험성평가 반영", f(9.5, True, NAVY), CENTER, HEAD); box(ws, f"E{r}:F{r}", None, f(10), LEFT)
    dv_list(ws, f"B{r}", ["TBM", "게시판", "안전교육", "산업안전보건위원회", "협의체 회의"]); dv_list(ws, f"E{r}", ["수시평가 반영", "반영 예정", "해당 없음"])
    ws.row_dimensions[r].height = 24
    legend(ws, r + 2, 6)
    ws.print_area = f"A1:F{r + 2}"

    lc = ["No", "발생일", "장소", "유형", "경위(요약)", "잠재 피해", "주요 원인", "재발방지 대책", "담당", "조치 기한", "완료일", "위험성평가\n반영", "상태"]
    m = sheet(wb, "아차사고 관리대장", [5, 11, 14, 12, 30, 10, 22, 30, 11, 11, 11, 11, 10], landscape=True, tab=ORANGE)
    title(m, "아차사고 관리대장", 9, sub="보고서를 한 줄로 모아 유형별 추세와 조치 완료 여부를 관리합니다")
    approval(m, 1, 11, labels=("담당", "승인"), sign_h=28)
    legend(m, 4, len(lc))
    header_row(m, 5, lc, height=30)
    f0, f1 = 6, 39
    grid(m, f0, f1, 1, len(lc), height=25)
    for x in range(f0, f1 + 1):
        m.cell(row=x, column=1, value=x - f0 + 1)
        for cc in (5, 7, 8):
            m.cell(row=x, column=cc).alignment = LEFT
        for cc in (2, 10, 11):
            m.cell(row=x, column=cc).number_format = "yyyy-mm-dd"
        m.cell(row=x, column=13, value=f'=IF(E{x}="","",IF(K{x}<>"","완료",IF(J{x}="","기한 미정",IF(J{x}<TODAY(),"기한 경과","진행 중"))))')
        m.cell(row=x, column=13).fill = fill(SUB)
    example(m, f0, [None, "2026-10-27", "자재창고 랙 3열", "물체에 맞음", "하역 중 상단 박스 낙하, 부상 없음", "중상 가능", "적재 높이 기준 없음", "적재 높이 1.5m 제한·출입 통제선", "물류팀", "2026-11-03", None, "수시평가 반영"])
    dv_list(m, f"D{f0}:D{f1}", TYPES); dv_list(m, f"F{f0}:F{f1}", ["사망 가능", "중상 가능", "경상 가능", "재산 피해"])
    dv_list(m, f"L{f0}:L{f1}", ["수시평가 반영", "반영 예정", "해당 없음"])
    for c in ("B", "J", "K"):
        dv_date(m, f"{c}{f0}:{c}{f1}")
    color_text_rule(m, f"M{f0}:M{f1}", "기한 경과", RED)
    color_text_rule(m, f"M{f0}:M{f1}", "완료", "D1FADF", "05603A")
    color_text_rule(m, f"F{f0}:F{f1}", "사망", RED)
    # 유형별 집계
    sr = f1 + 2
    m.cell(row=sr, column=2, value="유형별 건수").font = f(10.5, True, NAVY)
    for i, tname in enumerate(TYPES):
        rr = sr + 1 + i // 6; c0 = 2 + (i % 6) * 2
        box(m, f"{m.cell(row=rr, column=c0).column_letter}{rr}", tname, f(9, True, NAVY), CENTER, HEAD)
        box(m, f"{m.cell(row=rr, column=c0 + 1).column_letter}{rr}", f'=COUNTIF($D${f0}:$D${f1},"{tname}")', f(10, True), CENTER)
    m.print_title_rows = "5:5"; m.freeze_panes = "C6"; m.print_area = f"A1:M{sr + 2}"

    guide_sheet(wb, "아차사고 보고서·관리대장", [
        ("왜 기록하나", ["다친 사람이 없어도 같은 원인이 반복되면 사고로 이어집니다. 원인을 인적·물적·관리적으로 나눠 보면 대책이 구체적으로 나옵니다.",
                     "대책은 원인과 짝이 맞게 적고, 위험성평가(수시평가)에 반영했는지 표시하세요."]),
        ("주의", ["실명·연락처 같은 개인정보는 보고 범위에 맞게 최소한으로 적습니다. 보고한 사람에게 불이익이 없도록 운영해야 보고가 늘어납니다.",
                "실제 부상이 생겼다면 산업재해 조사·보고 대상인지 확인하세요(사망 또는 3일 이상 휴업: 산업재해조사표, 시행규칙 제73조)."]),
    ], source="아차사고 보고는 법정 서식이 없는 실무 기록입니다. 화면으로 작성: safeplum.com/tools/forms/accident/")
    wb.move_sheet("작성 안내", offset=-2); wb.active = 1
    p = OUT / "SafePlum_near_miss.xlsx"; wb.save(p); return p


# =====================================================================
# 6. 보호구 지급·관리대장
# =====================================================================
def ppe():
    law = json.load(open(ROOT / "data/lawtext/krule.json", encoding="utf-8"))
    art = next(a for a in law["articles"] if a.get("jo") == "제32조")
    import re
    lines = [ln for ln in art["text"].split("\n") if re.match(r"^[0-9]+\. ", ln)]
    wb = new_book()
    lc = ["No", "지급일", "소속", "성명", "보호구 종류", "규격·모델", "수량", "지급 사유", "사용기한·\n교체 예정일", "교체 상태", "수령 서명"]
    ws = sheet(wb, "보호구 지급대장", [5, 11, 13, 10, 14, 16, 6, 12, 12, 10, 14], landscape=True, tab=NAVY)
    title(ws, "보호구 지급·관리대장", 8, sub="작업조건에 맞는 보호구를 작업하는 근로자 수 이상 지급하고 착용하도록 해야 합니다(안전보건규칙 제32조)")
    approval(ws, 1, 9, labels=("담당", "승인"), sign_h=28)
    legend(ws, 4, len(lc))
    header_row(ws, 5, lc, height=30)
    f0, f1 = 6, 39
    grid(ws, f0, f1, 1, len(lc), height=24)
    for x in range(f0, f1 + 1):
        ws.cell(row=x, column=1, value=x - f0 + 1)
        for cc in (2, 9):
            ws.cell(row=x, column=cc).number_format = "yyyy-mm-dd"
        ws.cell(row=x, column=10, value=f'=IF(I{x}="","",IF(I{x}<TODAY(),"교체 필요",IF(I{x}-TODAY()<=30,"30일 이내","사용 중")))')
        ws.cell(row=x, column=10).fill = fill(SUB)
    KINDS = ["안전모", "안전대", "안전화", "보안경", "보안면", "절연용 보호구", "방열복", "방진마스크", "방독마스크", "송기마스크", "귀마개·귀덮개", "보호장갑", "방한복", "기타"]
    dv_list(ws, f"E{f0}:E{f1}", KINDS); dv_list(ws, f"H{f0}:H{f1}", ["신규 입사", "정기 교체", "파손·분실", "작업 변경"])
    for c in ("B", "I"):
        dv_date(ws, f"{c}{f0}:{c}{f1}")
    color_text_rule(ws, f"J{f0}:J{f1}", "교체 필요", RED)
    color_text_rule(ws, f"J{f0}:J{f1}", "30일 이내", ORANGE)
    example(ws, f0, [None, "2026-10-27", "생산1팀", "○○○", "안전화", "KCs 인증 ○○ 270", 1, "신규 입사", "2027-10-26"])
    sr = f1 + 2
    for i, (lab, frm) in enumerate([("지급 건수", f'=COUNTA(E{f0}:E{f1})'), ("지급 수량", f'=SUM(G{f0}:G{f1})'), ("교체 필요", f'=COUNTIF(J{f0}:J{f1},"교체 필요")'), ("30일 이내", f'=COUNTIF(J{f0}:J{f1},"30일 이내")')]):
        box(ws, f"{'BDFH'[i]}{sr}", lab, f(9.5, True, NAVY), CENTER, HEAD)
        box(ws, f"{'CEGI'[i]}{sr}", frm, f(11, True, RED if lab == "교체 필요" else INK), CENTER)
    ws.print_title_rows = "5:5"; ws.freeze_panes = "E6"; ws.print_area = f"A1:K{sr}"

    # 점검 기록
    k = sheet(wb, "보호구 점검기록", [5, 11, 14, 14, 30, 10, 26, 12], landscape=True, tab=GREEN)
    title(k, "보호구 점검 기록", 6, sub="정기적으로 보호구 상태를 점검하고 불량은 바로 교체합니다")
    header_row(k, 4, ["No", "점검일", "보관 장소·부서", "보호구 종류", "점검 내용", "결과", "조치", "다음 점검일"], height=26)
    grid(k, 5, 34, 1, 8, height=24)
    for x in range(5, 35):
        k.cell(row=x, column=1, value=x - 4)
        for cc in (2, 8):
            k.cell(row=x, column=cc).number_format = "yyyy-mm-dd"
        for cc in (5, 7):
            k.cell(row=x, column=cc).alignment = LEFT
    dv_list(k, "D5:D34", KINDS); dv_list(k, "F5:F34", ["양호", "불량", "교체"])
    color_text_rule(k, "F5:F34", "불량", RED)
    k.conditional_formatting.add("H5:H34", FormulaRule(formula=['AND(H5<>"",H5<TODAY())'], fill=fill("FEE4E2"), font=Font(name=FONT, color=RED, bold=True)))
    example(k, 5, [None, "2026-10-27", "용접 작업장", "보안면", "필터 렌즈 긁힘·변색, 머리띠 조임 상태", "불량", "필터 렌즈 교체 3개", "2026-11-27"])
    k.print_title_rows = "4:4"; k.print_area = "A1:H34"

    # 지급 기준
    b = sheet(wb, "지급 기준(제32조)", [6, 64, 30], tab="F79009")
    title(b, "작업별 보호구 지급 기준", 3, sub="산업안전보건기준에 관한 규칙 제32조제1항 (" + str(law.get("effective", "")) + " 시행본 기준)")
    header_row(b, 4, ["호", "작업", "보호구"], height=24)
    r = 5
    for ln in lines:
        num, rest = ln.split(".", 1)
        work, _, gear = rest.strip().rpartition(":")
        box(b, f"A{r}", num, f(10, True), CENTER); box(b, f"B{r}", work.strip(), f(9.5), LEFT); box(b, f"C{r}", gear.strip(), f(9.5, True, NAVY), LEFT)
        b.row_dimensions[r].height = 34 if len(work) > 50 else 24
        r += 1
    b.merge_cells(f"A{r + 1}:C{r + 1}")
    c = b.cell(row=r + 1, column=1, value="② 보호구를 받거나 착용지시를 받은 근로자는 그 보호구를 착용해야 합니다. — 사업장 작업별 지급 기준(소음 작업의 귀마개 등 다른 조문·작업환경 기준)은 아래 줄에 덧붙여 관리하세요.")
    c.font = f(9, False, MUTED); c.alignment = LEFT; b.row_dimensions[r + 1].height = 30
    b.print_area = f"A1:C{r + 1}"

    guide_sheet(wb, "보호구 지급·관리대장", [
        ("이렇게 쓰세요", ["지급할 때마다 한 줄씩 적고 수령 서명을 받습니다. 사용기한·교체 예정일을 넣으면 교체 상태가 자동으로 표시됩니다(지난 날짜: 교체 필요, 30일 이내: 주황).",
                       "보호구는 안전인증·자율안전확인 대상인지 확인하고 인증 제품을 지급하세요. 인증 확인: safeplum.com/tools/gear-cert/",
                       "점검기록 시트에 정기 점검 결과를 남기고, 다음 점검일이 지나면 빨간색으로 표시됩니다."]),
    ], source="근거: 산업안전보건기준에 관한 규칙 제32조(보호구의 지급 등). 화면으로 작성: safeplum.com/tools/forms/ppe-ledger/")
    wb.move_sheet("작성 안내", offset=-3); wb.active = 1
    p = OUT / "SafePlum_ppe_ledger.xlsx"; wb.save(p); return p


# =====================================================================
# 7. 작업계획서 (지게차·중량물·고소작업대·굴착기) — 포함 사항은 별표 4 원문 그대로
# =====================================================================
WP_SPECS = [
    {"sheet": "지게차", "no": "2", "id": "wp-forklift", "tab": "F79009", "lead": True,
     "name": "지게차 작업계획서", "law": "안전보건규칙 제38조제1항제2호 · 별표 4 제2호 · 제39조(작업지휘자)",
     "equip": [("장비 종류·번호", "예: 전동 지게차 3호기"), ("최대하중", "예: 2.5톤"), ("운전자·자격", "예: ○○○ / 면허·교육 이수 확인"), ("작업 시작 전 점검", "예: 별표 3 제7호 점검표 작성")],
     "pre": [("작업장소 노면·경사·통로 폭", "예: 창고 바닥 평탄, 출입구 경사 3% 구간 있음"), ("보행자 동선·주변 작업", "예: 포장 라인 작업자 통행로와 교차")],
     "plan_ex": ["예: 적재 높이·편하중 확인, 포크를 올린 채 주행 금지, 보행 통로 분리·교차로 일단정지, 유도자 배치, 하역 중 아래 출입 통제",
                 "예: 자재창고 → 2라인 입구(약 60m) 지정 경로, 사업장 제한속도 준수, 후진 시 경광등·후방 확인 / 팔레트 1단씩 하역"],
     "steps": [("상차·하역", "포크 위 화물이 떨어져 맞음", "화물 결속·랩핑 확인, 하역 반경 출입 통제"), ("운행", "보행자와 부딪힘", "지정 경로 운행, 교차로 일단정지, 유도자 배치")]},
    {"sheet": "중량물", "no": "11", "id": "wp-heavy", "tab": RED, "lead": True,
     "name": "중량물 취급 작업계획서", "law": "안전보건규칙 제38조제1항제11호 · 별표 4 제11호 · 제39조(작업지휘자) · 제385조",
     "equip": [("중량물 종류·형상", "예: 금형(상형) 1개, 직육면체"), ("무게·무게중심", "예: 약 1.8톤 / 중심 표시 있음"), ("취급 방법", "예: 천장크레인 2.8톤 + 슬링벨트 2줄 걸이"), ("신호수·유도", "예: 신호수 1명, 유도로프 1줄")],
     "pre": [("인양·운반 경로와 바닥 상태", "예: 금형 보관대 → 프레스 3호기, 경로 위 장애물 없음"), ("주변 작업·출입 통제 범위", "예: 인양 반경 3m 출입 통제")],
     "plan_ex": ["예: 작업대·적재함 위 작업 시 승강설비 사용, 끝단 접근 금지",
                 "예: 줄걸이 용구 작업 전 점검, 인양물 아래 출입 금지, 지면에서 약간 들어 균형 확인 후 인양",
                 "예: 무게중심 확인, 받침목·고임목 사용, 내려놓는 바닥 평탄 확인",
                 "예: 손이 끼이지 않게 받침목 위에 내림, 신호수 1명 지정, 유도로프 사용",
                 "예: 적재 단수 제한, 쌓은 물건 결속 상태 확인"],
     "steps": [("줄걸이", "슬링이 끊어져 인양물이 떨어짐", "작업 전 슬링 손상 점검, 정격하중 이내 사용"), ("내려놓기", "인양물과 금형대 사이에 손이 끼임", "받침목 사용, 손 대신 유도로프 사용")]},
    {"sheet": "고소작업대", "no": "2", "id": "wp-aerial", "tab": "2E90FA", "lead": True,
     "name": "고소작업대 작업계획서", "law": "안전보건규칙 제38조제1항제2호 · 별표 4 제2호 · 제186조",
     "equip": [("장비 종류·번호", "예: 시저형 고소작업대 1호기"), ("작업대 정격하중·최대 높이", "예: 230kg / 8m"), ("탑승 인원·운전자", "예: 2명 / ○○○ (조작 교육 이수)"), ("작업 시작 전 점검", "예: 별표 3 제11호 점검표 작성")],
     "pre": [("바닥 지지력·경사·개구부", "예: 1층 바닥 평탄, 배수구 덮개 위 설치 금지"), ("상부 장애물·전선", "예: 천장 배관·조명, 가공전선 없음")],
     "plan_ex": ["예: 작업대 안 안전대 체결, 난간 위 올라서기 금지, 공구 낙하 방지 끈, 작업대 아래 출입 통제",
                 "예: 작업 위치까지 낮춘 상태로 이동, 아웃트리거·브레이크 고정 후 상승, 작업대 높인 채 이동 시 유도자 배치"],
     "steps": [("상승·작업", "작업대에서 떨어짐", "안전대 체결, 난간 위 작업 금지"), ("이동", "상승 상태 이동 중 넘어짐", "낮춘 뒤 이동, 경사·턱 확인")]},
    {"sheet": "굴착기", "no": "3", "id": "wp-excavator", "tab": GREEN, "lead": False,
     "name": "차량계 건설기계(굴착기) 작업계획서", "law": "안전보건규칙 제38조제1항제3호 · 별표 4 제3호 · 제199조~제203조",
     "equip": [("장비 종류·규격", "예: 굴착기 0.6㎥(버킷)"), ("등록번호·검사", "예: 정기검사 유효기간 확인"), ("운전자(면허)", "예: ○○○ / 건설기계조종사면허 확인"), ("작업 시작 전 점검", "예: 브레이크·클러치·조종장치 기능")],
     "pre": [],
     "plan_ex": ["예: 굴착기 0.6㎥ 1대, 최대 굴착 깊이·작업 반경 확인", "예: 진입로(가설도로) → 굴착 구간, 갓길 붕괴 방지용 표지·유도자", "예: 버킷 인양 작업 금지(예외 요건 확인), 작업 반경 출입 통제, 운전석 이탈 시 버킷 지면에 내림"],
     "steps": [("굴착", "작업 반경 안 근로자와 부딪힘", "출입 통제선·유도자 배치"), ("이동", "갓길·경사에서 넘어짐", "지반 확인, 유도자 신호에 따라 이동")]},
]


def workplan():
    b4 = {it["no"]: it for it in json.load(open(ROOT / "data/byl4.json", encoding="utf-8"))["items"]}
    wb = new_book()
    W = [13, 13, 13, 12, 12, 13, 13, 13]
    for sp in WP_SPECS:
        it = b4[sp["no"]]
        ws = sheet(wb, sp["sheet"], W, tab=sp["tab"], one_page=True)
        title(ws, sp["name"], 4, sub=sp["law"])
        approval(ws, 1, 5, labels=("작성", "작업지휘자" if sp["lead"] else "책임자", "승인"), sign_h=30)
        hd = lambda ref, v: box(ws, ref, v, f(9.5, True, NAVY), CENTER, HEAD)
        sec = lambda r, v: box(ws, f"A{r}:H{r}", v, f(10.5, True, "FFFFFF"), LEFT, NAVY)
        # 1. 작업 개요
        r = 4
        sec(r, "1. 작업 개요"); r += 1
        lead_lab = "작업지휘자\n(제39조)" if sp["lead"] else "작업 책임자"
        for a, b_ in (("작업명", "작업 일시"), ("작업 장소", "작업 인원"), (lead_lab, "협력업체")):
            hd(f"A{r}", a); box(ws, f"B{r}:D{r}", None, f(10), LEFT)
            hd(f"E{r}", b_); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
            ws.row_dimensions[r].height = 26; r += 1
        example(ws, r - 3, [None, "예: " + {"wp-forklift": "원자재 입고 하역", "wp-heavy": "프레스 3호기 금형 교체", "wp-aerial": "1층 천장 조명 교체", "wp-excavator": "부지 배수관 매설 굴착"}[sp["id"]]])
        example(ws, r - 3, [None, None, None, None, None, "예: 2026-10-27 09:00~12:00"])
        # 2. 장비·대상
        r += 1; sec(r, "2. 사용 장비·취급 대상"); r += 1
        for i in range(0, len(sp["equip"]), 2):
            (a, ea), (b_, eb) = sp["equip"][i], sp["equip"][i + 1]
            hd(f"A{r}", a); box(ws, f"B{r}:D{r}", None, f(10), LEFT)
            hd(f"E{r}", b_); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
            example(ws, r, [None, ea, None, None, None, eb]); ws.row_dimensions[r].height = 30; r += 1
        # 3. 사전조사
        r += 1; sec(r, "3. 사전조사 결과 (제38조제1항 — 조사 결과를 기록·보존)"); r += 1
        header_row(ws, r, ["조사 항목", "", "", "조사 결과", "", "", "조치 필요 사항", ""], height=22)
        ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:F{r}"); ws.merge_cells(f"G{r}:H{r}"); r += 1
        pre_rows = [(x, None) for x in it["pre"]] + sp["pre"]
        for lab, ex in pre_rows:
            box(ws, f"A{r}:C{r}", lab, f(9.5, True, INK), LEFT, SUB); box(ws, f"D{r}:F{r}", None, f(10), LEFT); box(ws, f"G{r}:H{r}", None, f(10), LEFT)
            if ex:
                example(ws, r, [None, None, None, ex])
            ws.row_dimensions[r].height = 30 if len(lab) < 40 else 44; r += 1
        # 4. 작업계획 — 별표 4 포함 사항
        r += 1; sec(r, f"4. 작업계획 (별표 4 제{sp['no']}호에 따라 반드시 포함할 사항)"); r += 1
        header_row(ws, r, ["포함 사항(별표 4 원문)", "", "", "계획 내용", "", "", "", ""], height=22)
        ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:H{r}"); r += 1
        for i, lab in enumerate(it["plan"]):
            box(ws, f"A{r}:C{r}", lab, f(9.5, True, INK), LEFT, SUB); box(ws, f"D{r}:H{r}", None, f(10), LEFT_TOP)
            if i < len(sp["plan_ex"]):
                example(ws, r, [None, None, None, sp["plan_ex"][i]])
            ws.row_dimensions[r].height = 38; r += 1
        box(ws, f"A{r}:C{r}", "운행·인양 경로 약도\n(그림·사진 붙이기)", f(9.5, True, INK), LEFT, SUB); box(ws, f"D{r}:H{r}", None, f(10), LEFT_TOP)
        ws.row_dimensions[r].height = 72; r += 1
        # 5. 작업 단계별 위험요인
        r += 1; sec(r, "5. 작업 단계별 위험요인과 대책"); r += 1
        header_row(ws, r, ["작업 단계", "유해·위험요인", "", "", "안전대책", "", "", "담당"], height=22)
        ws.merge_cells(f"B{r}:D{r}"); ws.merge_cells(f"E{r}:G{r}"); r += 1
        s0 = r
        for k in range(4):
            grid(ws, r, r, 1, 8, height=26)
            ws.merge_cells(f"B{r}:D{r}"); ws.merge_cells(f"E{r}:G{r}")
            for cc in (2, 5):
                ws.cell(row=r, column=cc).alignment = LEFT
            if k < len(sp["steps"]):
                a, b_, c_ = sp["steps"][k]
                example(ws, r, [a, b_, None, None, c_])
            r += 1
        # 6. 비상 조치
        r += 1; sec(r, "6. 이상·비상 시 조치와 연락"); r += 1
        box(ws, f"A{r}:H{r}", None, f(10), LEFT_TOP); ws.row_dimensions[r].height = 40
        example(ws, r, ["예: 이상 발견 시 즉시 작업 중지 → 작업지휘자 보고 / 부상자 발생 시 119·사내 비상연락(내선 ○○○), 응급처치함 위치: ○○"]); r += 2
        # 7. 근로자 주지
        sec(r, "7. 작업계획서 내용 주지 확인 (제38조제2항 — 작업하는 근로자에게 알림)"); r += 1
        header_row(ws, r, ["소속", "성명", "서명", "", "소속", "성명", "서명", ""], height=22)
        ws.merge_cells(f"C{r}:D{r}"); ws.merge_cells(f"G{r}:H{r}"); r += 1
        for k in range(5):
            grid(ws, r, r, 1, 8, height=24)
            ws.merge_cells(f"C{r}:D{r}"); ws.merge_cells(f"G{r}:H{r}"); r += 1
        legend(ws, r + 1, 8)
        ws.print_area = f"A1:H{r + 1}"

    guide_sheet(wb, "작업계획서 (지게차·중량물·고소작업대·굴착기)", [
        ("언제 쓰나", ["안전보건규칙 제38조제1항 각 호의 작업(타워크레인 설치·해체, 차량계 하역운반기계등, 차량계 건설기계, 화학설비, 전기작업, 굴착, 터널, 교량, 채석, 해체, 중량물 취급, 궤도 보수·점검, 입환)을 할 때는 사전조사를 하고 그 결과를 기록·보존한 뒤, 별표 4의 사항을 포함한 작업계획서를 작성하고 그 계획에 따라 작업하도록 해야 합니다.",
                     "이 파일에는 현장에서 자주 쓰는 4종(지게차·중량물·고소작업대·굴착기)이 시트별로 들어 있습니다. 나머지 작업은 safeplum.com/tools/forms/workplan/ 에서 작업 종류를 골라 화면으로 작성하세요."]),
        ("이렇게 쓰세요", ["'4. 작업계획'의 왼쪽 회색 칸은 별표 4 원문 그대로이므로 지우지 말고, 오른쪽에 이 작업에 맞는 계획을 구체적으로 적습니다(장소·장비·인원·순서).",
                       "회색 기울임 글씨는 예시입니다. 지우고 사업장 내용으로 고쳐 쓰세요.",
                       "작성한 계획서 내용은 해당 근로자에게 알리고(제38조제2항) 서명을 받아 두세요. TBM에서 설명하고 서명받는 방법이 흔히 쓰입니다."]),
        ("작업지휘자", ["제38조제1항제2호(차량계 하역운반기계등)·제6호·제8호·제10호·제11호(중량물 취급) 작업계획서를 작성한 경우 작업지휘자를 지정해 계획서에 따라 작업을 지휘하도록 해야 합니다(제39조제1항). 제2호 작업은 다른 근로자가 접근할 수 없거나 주위에 근로자가 없어 충돌 위험이 없는 1대 운전 작업이면 지정하지 않을 수 있습니다(같은 항 단서)."]),
    ], source="근거: 산업안전보건기준에 관한 규칙 제38조·제39조, 별표 4 (" + json.load(open(ROOT / "data/byl4.json", encoding="utf-8")).get("checked", "") + " 원문 확인본). 법령 서식이 따로 정해져 있지 않아 사업장 양식으로 쓸 수 있으며, 별표 4 사항이 모두 들어가야 합니다.")
    wb.move_sheet("작성 안내", offset=-len(WP_SPECS)); wb.active = 1
    p = OUT / "SafePlum_work_plan.xlsx"; wb.save(p); return p


# =====================================================================
# 8. 밀폐공간 — 작업 전 확인서(제619조제2항)·측정 기록(제619조의2)·출입 인원(제621조)·관리대장(제619조제1항, 제640조)
# =====================================================================
GAS_OK = "적정공기: 산소 18% 이상 23.5% 미만 · 이산화탄소 1.5% 미만 · 일산화탄소 30ppm 미만 · 황화수소 10ppm 미만 (제618조제3호)"


def _gas_judge(r, o2, co2, co, h2s):
    return (f'=IF(COUNT({o2}{r},{co2}{r},{co}{r},{h2s}{r})=0,"",IF(COUNT({o2}{r},{co2}{r},{co}{r},{h2s}{r})<4,"미측정 항목",'
            f'IF(AND({o2}{r}>=18,{o2}{r}<23.5,{co2}{r}<1.5,{co}{r}<30,{h2s}{r}<10),"적정","부적정 — 작업 금지")))')


def confined():
    wb = new_book()
    W = [12, 12, 11, 11, 11, 11, 13, 13]
    ws = sheet(wb, "작업 전 확인서", W, tab=NAVY, one_page=True)
    title(ws, "밀폐공간 작업 전 확인서", 4, sub="안전보건규칙 제619조제2항 — 작업이 끝날 때까지 출입구에 게시(제3항)")
    approval(ws, 1, 5, labels=("관리감독자", "안전관리자", "허가"), sign_h=30)
    hd = lambda ref, v: box(ws, ref, v, f(9.5, True, NAVY), CENTER, HEAD)
    sec = lambda r, v: box(ws, f"A{r}:H{r}", v, f(10.5, True, "FFFFFF"), LEFT, NAVY)
    r = 4; sec(r, "1. 작업 정보 (일시·기간·장소·내용)"); r += 1
    for a, b_, ea, eb in (("작업 일시", "작업 기간", "예: 2026-10-27 09:00~11:30", "예: 1일"), ("밀폐공간", "작업 인원", "예: 오폐수처리장 집수정 2호", "예: 3명(작업 2·감시 1)")):
        hd(f"A{r}", a); box(ws, f"B{r}:D{r}", None, f(10), LEFT); hd(f"E{r}", b_); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
        example(ws, r, [None, ea, None, None, None, eb]); ws.row_dimensions[r].height = 24; r += 1
    hd(f"A{r}", "작업 내용"); box(ws, f"B{r}:H{r}", None, f(10), LEFT_TOP); example(ws, r, [None, "예: 집수정 바닥 슬러지 제거 및 수중펌프 교체"]); ws.row_dimensions[r].height = 30; r += 2
    sec(r, "2. 작업자 정보 (관리감독자·감시인·측정자·작업자)"); r += 1
    header_row(ws, r, ["역할", "성명", "소속", "", "연락처", "", "교육·숙지 확인", ""], height=22)
    for a, b_ in (("C", "D"), ("E", "F"), ("G", "H")):
        ws.merge_cells(f"{a}{r}:{b_}{r}")
    r += 1
    for role in ("관리감독자", "감시인", "측정자", "작업자", "작업자", "작업자"):
        grid(ws, r, r, 1, 8, height=22)
        for a, b_ in (("C", "D"), ("E", "F"), ("G", "H")):
            ws.merge_cells(f"{a}{r}:{b_}{r}")
        c = ws.cell(row=r, column=1, value=role); c.font = f(9.5, True, INK); c.fill = fill(SUB)
        r += 1
    dv_list(ws, f"G{r - 6}:G{r - 1}", ["확인", "미확인"])
    r += 1
    sec(r, "3. 산소·유해가스 측정 결과와 후속조치 (작업 시작·재개 전마다, 제619조의2)"); r += 1
    header_row(ws, r, ["측정 시각", "위치", "산소 %", "이산화탄소 %", "일산화탄소 ppm", "황화수소 ppm", "판정(자동)", "후속 조치"], height=30); r += 1
    g0 = r
    for k in range(4):
        grid(ws, r, r, 1, 8, height=22)
        ws.cell(row=r, column=7, value=_gas_judge(r, "C", "D", "E", "F")).fill = fill(SUB)
        ws.cell(row=r, column=1).number_format = "hh:mm"
        r += 1
    example(ws, g0, ["09:05", "상·중·하", 20.8, 0.1, 2, 0, None, "환기 유지"])
    dv_list(ws, f"B{g0}:B{r - 1}", ["상", "중", "하", "상·중·하"])
    color_text_rule(ws, f"G{g0}:G{r - 1}", "부적정", RED)
    color_text_rule(ws, f"G{g0}:G{r - 1}", "미측정", ORANGE, INK)
    ws.conditional_formatting.add(f"G{g0}:G{r - 1}", FormulaRule(formula=[f'G{g0}="적정"'], fill=fill("D1FADF"), font=Font(name=FONT, color="05603A", bold=True)))
    box(ws, f"A{r}:H{r}", GAS_OK, f(8.5, False, MUTED), LEFT); ws.row_dimensions[r].height = 18; r += 2
    sec(r, "4. 작업 중 불활성가스·유해가스 누출·유입·발생 가능성 검토와 후속조치"); r += 1
    header_row(ws, r, ["검토 항목", "", "", "가능성", "후속 조치", "", "", ""], height=22)
    ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"E{r}:H{r}"); r += 1
    c0 = r
    for lab, ex in (("연결 배관·밸브로 유입", "예: 유입 밸브 잠금·표지, 맹판 설치"), ("불활성가스(질소 등) 퍼지·누출", None),
                    ("슬러지·부패물·화학반응으로 발생", "예: 교반 시 황화수소 발생 — 연속 측정, 송기마스크"), ("작업 자체에서 발생(용접·도장 등)", None)):
        box(ws, f"A{r}:C{r}", lab, f(9.5, False, INK), LEFT, SUB); box(ws, f"D{r}", None, f(10), CENTER); box(ws, f"E{r}:H{r}", None, f(10), LEFT)
        if ex:
            example(ws, r, [None, None, None, "있음", ex])
        ws.row_dimensions[r].height = 24; r += 1
    dv_list(ws, f"D{c0}:D{r - 1}", ["있음", "없음", "검토 중"])
    color_text_rule(ws, f"D{c0}:D{r - 1}", "있음", ORANGE, INK)
    r += 1
    sec(r, "5. 착용할 보호구 · 환기 (제620조)"); r += 1
    gear = ["공기호흡기", "송기마스크", "안전대·구명밧줄", "안전모", "휴대용 가스측정기", "기타"]
    for i in range(0, 6, 3):
        for j in range(3):
            col = 1 + j * 3 if j < 2 else 7
            a = ws.cell(row=r, column=col)
            box(ws, a.coordinate, None, f(11, True, GREEN), CENTER)
            last = col + 2 if j < 2 else 8
            box(ws, f"{ws.cell(row=r, column=col + 1).coordinate}:{ws.cell(row=r, column=last).coordinate}", gear[i + j], f(9.5), LEFT)
        dv_list(ws, f"A{r}", ["✔"]); dv_list(ws, f"D{r}", ["✔"]); dv_list(ws, f"G{r}", ["✔"])
        ws.row_dimensions[r].height = 22; r += 1
    hd(f"A{r}", "환기 방법"); box(ws, f"B{r}:H{r}", None, f(10), LEFT); example(ws, r, [None, "예: 송풍기 1대 작업 30분 전부터 작업 끝까지 연속 가동"]); ws.row_dimensions[r].height = 24; r += 2
    sec(r, "6. 비상연락체계"); r += 1
    for a, b_, ea, eb in (("감시인 → 보고", "구조 장비 위치", "예: 감시인 → 관리감독자(010-…) → 119·사내 비상(내선 …)", "예: 공기호흡기 2조·삼각대 — 처리장 입구"),):
        hd(f"A{r}", a); box(ws, f"B{r}:D{r}", None, f(10), LEFT); hd(f"E{r}", b_); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
        example(ws, r, [None, ea, None, None, None, eb]); ws.row_dimensions[r].height = 34; r += 1
    box(ws, f"A{r}:H{r}", "사고가 나면 감시인은 들어가지 말고 구조 요청 — 구조하는 사람도 공기호흡기 또는 송기마스크를 착용해야 합니다.", f(9, True, RED), LEFT)
    ws.row_dimensions[r].height = 20; r += 1
    legend(ws, r + 1, 8)
    ws.print_area = f"A1:H{r + 1}"

    # 측정 기록 (3년 보존)
    m = sheet(wb, "측정 기록(3년 보존)", [5, 15, 18, 11, 9, 9, 10, 10, 10, 15, 22, 11], landscape=True, tab=GREEN)
    title(m, "밀폐공간 산소·유해가스 측정 기록", 9, sub="측정·평가한 사람의 성명, 일시·장소, 결과를 기록해 3년간 보존 (안전보건규칙 제619조의2제4항)")
    lc = ["No", "측정 일시", "장소(밀폐공간)", "시점", "위치", "산소 %", "이산화탄소 %", "일산화탄소\nppm", "황화수소\nppm", "판정(자동)", "조치", "측정자 성명"]
    legend(m, 3, len(lc), text=GAS_OK)
    header_row(m, 4, lc, height=30)
    f0, f1 = 5, 44
    grid(m, f0, f1, 1, len(lc), height=22)
    for x in range(f0, f1 + 1):
        m.cell(row=x, column=1, value=x - f0 + 1)
        m.cell(row=x, column=2).number_format = "yyyy-mm-dd hh:mm"
        m.cell(row=x, column=10, value=_gas_judge(x, "F", "G", "H", "I")).fill = fill(SUB)
        m.cell(row=x, column=11).alignment = LEFT
    example(m, f0, [None, "2026-10-27 09:05", "집수정 2호", "작업 전", "상·중·하", 20.8, 0.1, 2, 0, None, "환기 유지", "○○○"])
    dv_list(m, f"D{f0}:D{f1}", ["작업 전", "재개 전", "작업 중", "작업 후"])
    dv_list(m, f"E{f0}:E{f1}", ["상", "중", "하", "상·중·하"])
    color_text_rule(m, f"J{f0}:J{f1}", "부적정", RED)
    color_text_rule(m, f"J{f0}:J{f1}", "미측정", ORANGE, INK)
    m.print_title_rows = "4:4"; m.freeze_panes = "D5"; m.print_area = f"A1:L{f1}"

    # 출입 인원 점검
    io = sheet(wb, "출입 인원 점검", [5, 12, 14, 12, 12, 12, 14, 18], tab="2E90FA")
    title(io, "밀폐공간 출입 인원 점검표", 8, sub="입장시킬 때와 퇴장시킬 때마다 인원을 점검 (안전보건규칙 제621조)")
    hd2 = lambda ref, v: box(io, ref, v, f(9.5, True, NAVY), CENTER, HEAD)
    hd2("A4:B4", "일자·장소"); box(io, "C4:F4", None, f(10), LEFT); hd2("G4", "감시인"); box(io, "H4", None, f(10), LEFT)
    io.row_dimensions[4].height = 24
    header_row(io, 6, ["No", "성명", "소속", "입장 시각", "퇴장 시각", "체류(자동)", "산소측정 확인", "감시인 확인"], height=26)
    grid(io, 7, 31, 1, 8, height=22)
    for x in range(7, 32):
        io.cell(row=x, column=1, value=x - 6)
        for cc in (4, 5):
            io.cell(row=x, column=cc).number_format = "hh:mm"
        io.cell(row=x, column=6, value=f'=IF(OR(D{x}="",E{x}=""),IF(D{x}<>"","작업 중",""),E{x}-D{x})').number_format = "[h]:mm"
        io.cell(row=x, column=6).fill = fill(SUB)
    example(io, 7, [None, "○○○", "환경팀", "09:12", "09:58", None, "적정", "○"])
    dv_list(io, "G7:G31", ["적정", "미확인"])
    box(io, "A33:C33", "입장 인원", f(9.5, True, NAVY), CENTER, HEAD); box(io, "D33", "=COUNTA(D7:D31)", f(11, True), CENTER)
    box(io, "E33:F33", "아직 안에 있는 인원", f(9.5, True, NAVY), CENTER, HEAD); box(io, "G33", '=COUNTIF(F7:F31,"작업 중")', f(11, True, RED), CENTER)
    io.conditional_formatting.add("F7:F31", FormulaRule(formula=['F7="작업 중"'], fill=fill(YELLOW), font=Font(name=FONT, bold=True)))
    io.print_area = "A1:H33"

    # 관리대장 + 긴급구조훈련
    g = sheet(wb, "밀폐공간 관리대장", [5, 18, 18, 14, 22, 18, 12, 12, 16], landscape=True, tab=ORANGE, one_page=True)
    title(g, "밀폐공간 관리대장 · 긴급구조훈련 기록", 9, sub="밀폐공간 작업 프로그램의 '위치 파악 및 관리 방안'·'유해·위험 요인 파악'(제619조제1항) · 긴급구조훈련 6개월에 1회 이상(제640조)")
    header_row(g, 4, ["No", "밀폐공간명", "위치", "별표 18 해당", "예상 유해·위험 요인", "출입 통제 방법", "관리 부서", "최근 점검일", "비고"], height=28)
    grid(g, 5, 19, 1, 9, height=22)
    for x in range(5, 20):
        g.cell(row=x, column=1, value=x - 4); g.cell(row=x, column=8).number_format = "yyyy-mm-dd"
        for cc in (5, 6):
            g.cell(row=x, column=cc).alignment = LEFT
    example(g, 5, [None, "집수정 2호", "오폐수처리장 지하", "제○호(원문 확인)", "황화수소·산소결핍", "맨홀 잠금·출입금지 표지", "환경팀", "2026-10-20"])
    dv_list(g, "E5:E19", ["산소결핍", "황화수소", "일산화탄소", "이산화탄소", "가연성가스", "황화수소·산소결핍"])
    box(g, "A21:I21", "긴급구조훈련 기록 (비상연락체계 운영·구조용 장비 사용·공기호흡기 또는 송기마스크 착용·응급처치 — 6개월에 1회 이상, 결과 기록·보존)", f(10, True, "FFFFFF"), LEFT, NAVY)
    g.row_dimensions[21].height = 22
    header_row(g, 22, ["No", "훈련일", "대상 밀폐공간", "참석 인원", "훈련 내용", "보완 사항", "다음 훈련 기한(자동)", "", "확인"], height=26)
    g.merge_cells("G22:H22")
    grid(g, 23, 30, 1, 9, height=22)
    for x in range(23, 31):
        g.merge_cells(f"G{x}:H{x}")
        g.cell(row=x, column=1, value=x - 22); g.cell(row=x, column=2).number_format = "yyyy-mm-dd"
        g.cell(row=x, column=7, value=f'=IF(B{x}="","",EDATE(B{x},6))').number_format = "yyyy-mm-dd"
        g.cell(row=x, column=7).fill = fill(SUB)
        for cc in (5, 6):
            g.cell(row=x, column=cc).alignment = LEFT
    g.conditional_formatting.add("G23:G30", FormulaRule(formula=['AND(G23<>"",G23<TODAY(),COUNTIF($B$23:$B$30,">"&B23)=0)'], fill=fill("FEE4E2"), font=Font(name=FONT, color=RED, bold=True)))
    example(g, 23, [None, "2026-10-15", "집수정 2호", "8명", "공기호흡기 착용, 삼각대로 구조, 심폐소생술", "구조 로프 길이 부족 — 교체"])
    g.print_area = "A1:I30"

    guide_sheet(wb, "밀폐공간 작업 양식", [
        ("이 파일의 구성", ["작업 전 확인서: 제619조제2항 6가지(작업 정보·작업자 정보·산소/유해가스 측정 결과와 후속조치·누출 등 가능성 검토·보호구·비상연락체계)를 한 장에 확인하고, 작업이 끝날 때까지 출입구에 붙입니다(제3항).",
                       "측정 기록: 작업 시작 전(일시 중단 후 다시 시작할 때 포함) 지정된 측정자가 산소·유해가스를 측정·평가하고, 측정자 성명·일시·장소·결과를 기록해 3년간 보존합니다(제619조의2제1항·제4항). 숫자를 넣으면 적정공기 기준(제618조제3호)과 비교해 판정이 자동으로 나옵니다.",
                       "출입 인원 점검: 들어갈 때와 나올 때마다 인원을 점검합니다(제621조). 퇴장 시각이 비어 있으면 '작업 중'으로 표시됩니다.",
                       "관리대장: 사업장 밀폐공간의 위치·유해 요인·출입 통제를 정리하고(밀폐공간 작업 프로그램, 제619조제1항), 6개월에 1회 이상 긴급구조훈련 결과를 남깁니다(제640조)."]),
        ("주의", ["판정은 넣은 숫자를 기준과 비교한 결과일 뿐이며, 측정 위치·장비 교정 상태가 적절한지는 판단하지 않습니다. 네 가지 중 하나라도 비어 있으면 '미측정 항목'으로 표시됩니다.",
                "밀폐공간에 해당하는 장소는 안전보건규칙 별표 18 원문에서 확인해 '별표 18 해당' 칸에 적으세요.",
                "적정공기가 유지되지 않으면 환기하거나 공기호흡기·송기마스크를 지급해 착용시키는 등 필요한 조치를 해야 합니다(제619조의2제3항, 제620조)."]),
    ], source="근거: 산업안전보건기준에 관한 규칙 제618조·제619조·제619조의2·제620조·제621조·제640조. 화면으로 작성: safeplum.com/tools/confined/")
    wb.move_sheet("작성 안내", offset=-4); wb.active = 1
    p = OUT / "SafePlum_confined_space.xlsx"; wb.save(p); return p


# =====================================================================
# 9. 화기작업 허가서 — 제241조(화재위험작업 시의 준수사항)·제241조의2(화재감시자)
# =====================================================================
def hot_work():
    k = json.load(open(ROOT / "data/lawtext/krule.json", encoding="utf-8"))
    import re
    a241 = next(a for a in k["articles"] if a.get("jo") == "제241조")["text"]
    items = [re.sub(r"\. 이 경우.*$", "", ln.split(". ", 1)[1]) for ln in a241.split("\n") if re.match(r"^[1-6]\. ", ln)]
    assert len(items) == 6, items
    wb = new_book()
    W = [12, 12, 12, 12, 12, 12, 13, 13]
    ws = sheet(wb, "화기작업 허가서", W, tab=RED, one_page=True)
    title(ws, "화기작업 허가서", 4, sub="안전보건규칙 제241조·제241조의2 — 작업 시작부터 끝날 때까지 작업장소에 게시(제241조제4항)")
    approval(ws, 1, 5, labels=("신청", "안전관리자", "허가"), sign_h=30)
    hd = lambda ref, v: box(ws, ref, v, f(9.5, True, NAVY), CENTER, HEAD)
    sec = lambda r, v, c=NAVY: box(ws, f"A{r}:H{r}", v, f(10.5, True, "FFFFFF"), LEFT, c)
    r = 4; sec(r, "1. 작업 개요"); r += 1
    for a, b_, ea, eb in (("작업 일시", "작업 종류", "예: 2026-10-27 13:00~16:00", "용접"), ("작업 장소", "작업 업체", "예: 2층 기계실 배관 랙", "예: ○○설비(협력업체)"),
                          ("작업자", "화재감시자", "예: ○○○ 외 1명", "예: ○○○")):
        hd(f"A{r}", a); box(ws, f"B{r}:D{r}", None, f(10), LEFT); hd(f"E{r}", b_); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
        example(ws, r, [None, ea, None, None, None, eb]); ws.row_dimensions[r].height = 24; r += 1
    dv_list(ws, f"F{r - 3}", ["용접", "용단", "금속 가열", "연삭숫돌 건식연마", "그라인딩", "기타 불꽃 발생 작업"])
    hd(f"A{r}", "작업 내용"); box(ws, f"B{r}:H{r}", None, f(10), LEFT_TOP); example(ws, r, [None, "예: 냉각수 배관 플랜지 교체를 위한 용단·용접, 주변 보온재 있음"]); ws.row_dimensions[r].height = 30; r += 2
    sec(r, "2. 작업 전 확인 — 가연성물질이 있는 장소 (제241조제2항 각 호, 원문)"); r += 1
    header_row(ws, r, ["확인 사항(원문)", "", "", "", "확인", "확인 내용", "", ""], height=22)
    ws.merge_cells(f"A{r}:D{r}"); ws.merge_cells(f"F{r}:H{r}"); r += 1
    c0 = r
    ex = ["예: 작업 순서·담당 정해 TBM 공유", "예: 도료 18L 2통 → 창고로 이동", "예: 보온재 방화포로 덮음, 소화기 2대", "예: 성능인증 용접방화포 사용", "예: 가스 측정 0%LEL 확인", "예: 소화기 위치·대피로 교육"]
    for i, it in enumerate(items):
        box(ws, f"A{r}:D{r}", f"{i + 1}. {it}", f(9.5, False, INK), LEFT, SUB); box(ws, f"E{r}", None, f(10, True), CENTER); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
        example(ws, r, [None, None, None, None, None, ex[i]])
        ws.row_dimensions[r].height = 30 if len(it) < 34 else 40; r += 1
    box(ws, f"A{r}:D{r}", "통풍·환기가 충분하지 않은 장소에서 통풍·환기용으로 산소를 쓰지 않음 (제1항)", f(9.5, False, INK), LEFT, SUB); box(ws, f"E{r}", None, f(10, True), CENTER); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
    ws.row_dimensions[r].height = 30; r += 1
    dv_list(ws, f"E{c0}:E{r - 1}", ["확인", "해당 없음", "미확인"])
    color_text_rule(ws, f"E{c0}:E{r - 1}", "미확인", RED)
    r += 1
    sec(r, "3. 화재감시자 배치 (제241조의2) — 용접·용단 작업"); r += 1
    header_row(ws, r, ["배치 대상 장소(어느 하나에 해당하면 배치)", "", "", "", "", "", "해당 여부", ""], height=22)
    ws.merge_cells(f"A{r}:F{r}"); ws.merge_cells(f"G{r}:H{r}"); r += 1
    k0 = r
    for t in ("작업반경 11미터 이내에 건물구조 자체나 내부(개구부 등으로 개방된 부분 포함)에 가연성물질이 있는 장소",
              "작업반경 11미터 이내의 바닥 하부에 가연성물질이 11미터 이상 떨어져 있지만 불꽃에 의해 쉽게 발화될 우려가 있는 장소",
              "가연성물질이 금속으로 된 칸막이·벽·천장 또는 지붕의 반대쪽 면에 인접해 있어 열전도나 열복사에 의해 발화될 우려가 있는 장소"):
        box(ws, f"A{r}:F{r}", t, f(9, False, INK), LEFT, SUB); box(ws, f"G{r}:H{r}", None, f(10, True), CENTER); ws.row_dimensions[r].height = 30; r += 1
    dv_list(ws, f"G{k0}:G{r - 1}", ["해당", "해당 없음"])
    hd(f"A{r}", "감시자 지급 장비"); box(ws, f"B{r}:H{r}", "□ 확성기   □ 휴대용 조명기구   □ 화재 대피용 마스크   (제241조의2제3항)", f(10), LEFT); ws.row_dimensions[r].height = 24; r += 1
    hd(f"A{r}", "감시자 업무"); box(ws, f"B{r}:H{r}", "가연성물질 유무 확인 · 가스 검지·경보 장치 작동 확인 · 화재 시 근로자 대피 유도 (제241조의2제2항)", f(9.5, False, MUTED), LEFT); ws.row_dimensions[r].height = 24; r += 2
    sec(r, "4. 인화성 가스·증기 측정 (작업 전·작업 중)"); r += 1
    header_row(ws, r, ["측정 시각", "위치", "", "가연성가스 %LEL", "", "측정자", "", "조치"], height=22)
    ws.merge_cells(f"B{r}:C{r}"); ws.merge_cells(f"D{r}:E{r}"); ws.merge_cells(f"F{r}:G{r}"); r += 1
    for k_ in range(3):
        grid(ws, r, r, 1, 8, height=22)
        ws.merge_cells(f"B{r}:C{r}"); ws.merge_cells(f"D{r}:E{r}"); ws.merge_cells(f"F{r}:G{r}"); r += 1
    example(ws, r - 3, ["12:50", "배관 랙 하부", None, "0", None, "○○○", None, "작업 가능"])
    r += 1
    sec(r, "5. 작업 종료 확인 — 불씨·잔불", RED); r += 1
    hd(f"A{r}", "종료 시각"); box(ws, f"B{r}:D{r}", None, f(10), LEFT); hd(f"E{r}", "잔불 확인 시각"); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
    example(ws, r, [None, "15:40", None, None, None, "예: 16:10 (종료 30분 뒤 순회)"]); ws.row_dimensions[r].height = 24; r += 1
    hd(f"A{r}", "확인자"); box(ws, f"B{r}:D{r}", None, f(10), LEFT); hd(f"E{r}", "특이 사항"); box(ws, f"F{r}:H{r}", None, f(10), LEFT); ws.row_dimensions[r].height = 24; r += 1
    legend(ws, r + 1, 8)
    ws.print_area = f"A1:H{r + 1}"

    lg = sheet(wb, "화기작업 대장", [5, 11, 9, 9, 18, 13, 22, 12, 12, 11, 11], landscape=True, tab=ORANGE)
    title(lg, "화기작업 허가 대장", 8, sub="허가서를 한 줄로 모아 허가 번호·장소·감시자·잔불 확인을 관리합니다")
    lc = ["No", "작업일", "시작", "종료", "장소", "작업 종류", "작업 내용", "작업 업체", "화재감시자", "잔불 확인", "허가자"]
    header_row(lg, 4, lc, height=26)
    grid(lg, 5, 44, 1, len(lc), height=22)
    for x in range(5, 45):
        lg.cell(row=x, column=1, value=x - 4); lg.cell(row=x, column=2).number_format = "yyyy-mm-dd"
        for cc in (3, 4):
            lg.cell(row=x, column=cc).number_format = "hh:mm"
        for cc in (5, 7):
            lg.cell(row=x, column=cc).alignment = LEFT
    dv_list(lg, "F5:F44", ["용접", "용단", "금속 가열", "연삭숫돌 건식연마", "그라인딩", "기타 불꽃 발생 작업"])
    dv_list(lg, "J5:J44", ["확인", "미확인"])
    color_text_rule(lg, "J5:J44", "미확인", RED)
    lg.conditional_formatting.add("J5:J44", FormulaRule(formula=['AND(B5<>"",J5="")'], fill=fill(YELLOW)))
    example(lg, 5, [None, "2026-10-27", "13:00", "15:40", "2층 기계실 배관 랙", "용접", "냉각수 배관 플랜지 교체", "○○설비", "○○○", "확인", "○○○"])
    lg.print_title_rows = "4:4"; lg.freeze_panes = "C5"; lg.print_area = "A1:K44"

    guide_sheet(wb, "화기작업 허가서", [
        ("언제 쓰나", ["용접·용단, 금속 가열, 연삭숫돌 건식연마 등 불꽃이 생길 수 있는 작업(화재위험작업, 제236조제2항)을 가연성물질이 있는 장소에서 할 때 씁니다.",
                     "작업 전에 제241조제2항 각 호를 확인하고 비산 방지 등 안전조치를 한 뒤 작업하게 해야 하며(제3항), 작업 내용·일시·안전점검 및 조치 사항을 작업이 끝날 때까지 작업장소에 서면으로 게시해야 합니다(제4항). 같은 장소에서 상시·반복적으로 하는 작업은 게시를 생략할 수 있습니다."]),
        ("화재감시자", ["용접·용단 작업 장소가 제241조의2제1항 각 호 중 하나에 해당하면 화재감시자를 지정해 배치합니다. 같은 장소에서 상시·반복적으로 하면서 경보용 설비·기구, 소화설비 또는 소화기를 갖춘 경우에는 배치하지 않을 수 있습니다(같은 항 단서).",
                     "감시자에게 확성기·휴대용 조명기구·화재 대피용 마스크 등 대피용 방연장비를 지급해야 합니다(제3항)."]),
        ("이렇게 쓰세요", ["'2. 작업 전 확인'의 회색 칸은 원문입니다. 확인 칸에서 확인·해당 없음·미확인을 고르면 미확인은 빨간색으로 표시됩니다.",
                       "작업이 끝난 뒤 잔불 확인 시각과 확인자를 반드시 적고, 대장에 옮겨 관리하세요. 잔불 확인 칸이 비어 있으면 노란색으로 표시됩니다."]),
    ], source="근거: 산업안전보건기준에 관한 규칙 제236조제2항·제241조·제241조의2. 사업장 안전작업허가 절차가 있으면 그 절차를 따르세요. 화면으로 작성: safeplum.com/tools/permit/")
    wb.move_sheet("작성 안내", offset=-2); wb.active = 1
    p = OUT / "SafePlum_hot_work_permit.xlsx"; wb.save(p); return p


# =====================================================================
# 10. 안전 및 보건에 관한 협의체 회의록 — 시행규칙 제79조
# =====================================================================
def council():
    r79 = next(a for a in json.load(open(ROOT / "data/lawtext/rule.json", encoding="utf-8"))["articles"] if a.get("jo") == "제79조")["text"]
    import re
    items = [ln.split(". ", 1)[1].replace("법 제36조에 따른 ", "") for ln in r79.split("\n") if re.match(r"^[1-5]\. ", ln)]
    assert len(items) == 5, items
    wb = new_book()
    W = [12, 12, 12, 12, 12, 12, 13, 13]
    ws = sheet(wb, "협의체 회의록", W, tab=NAVY, one_page=True)
    title(ws, "안전 및 보건에 관한 협의체 회의록", 4, sub="산업안전보건법 제64조제1항제1호 · 시행규칙 제79조 — 매월 1회 이상 정기 회의, 결과 기록·보존")
    approval(ws, 1, 5, labels=("작성", "안전관리자", "승인(도급인)"), sign_h=30)
    hd = lambda ref, v: box(ws, ref, v, f(9.5, True, NAVY), CENTER, HEAD)
    sec = lambda r, v: box(ws, f"A{r}:H{r}", v, f(10.5, True, "FFFFFF"), LEFT, NAVY)
    r = 4; sec(r, "1. 회의 개요"); r += 1
    for a, b_, ea, eb in (("회의 일시", "회차", "예: 2026-10-27 14:00~15:00", "예: 2026년 10월 정기"), ("장소", "작성자", "예: 본관 2층 회의실", "예: 안전관리자 ○○○")):
        hd(f"A{r}", a); box(ws, f"B{r}:D{r}", None, f(10), LEFT); hd(f"E{r}", b_); box(ws, f"F{r}:H{r}", None, f(10), LEFT)
        example(ws, r, [None, ea, None, None, None, eb]); ws.row_dimensions[r].height = 24; r += 1
    r += 1
    sec(r, "2. 참석자 — 도급인 및 그의 수급인 전원으로 구성 (제79조제1항)"); r += 1
    header_row(ws, r, ["구분", "업체명", "", "맡은 작업(공종)", "", "직위", "성명", "서명"], height=22)
    ws.merge_cells(f"B{r}:C{r}"); ws.merge_cells(f"D{r}:E{r}"); r += 1
    a0 = r
    for k_ in range(7):
        grid(ws, r, r, 1, 8, height=22)
        ws.merge_cells(f"B{r}:C{r}"); ws.merge_cells(f"D{r}:E{r}"); r += 1
    dv_list(ws, f"A{a0}:A{r - 1}", ["도급인", "수급인"])
    example(ws, a0, ["도급인", "(주)○○", None, "전체 관리", None, "공장장", "○○○"])
    example(ws, a0 + 1, ["수급인", "○○설비", None, "배관 보수", None, "현장소장", "○○○"])
    box(ws, f"A{r}:B{r}", "참석 업체 수(자동)", f(9.5, True, NAVY), CENTER, HEAD); box(ws, f"C{r}", f'=SUMPRODUCT((B{a0}:B{r - 1}<>"")/COUNTIF(B{a0}:B{r - 1},B{a0}:B{r - 1}&""))', f(11, True), CENTER)
    box(ws, f"D{r}:H{r}", "수급인이 모두 참석했는지 확인하세요. 불참 업체는 회의 결과를 따로 전달하고 기록합니다.", f(8.5, False, MUTED), LEFT)
    ws.row_dimensions[r].height = 20; r += 2
    sec(r, "3. 협의 사항 (제79조제2항 각 호 — 원문)"); r += 1
    header_row(ws, r, ["협의 사항", "", "", "협의 내용·결정", "", "", "", ""], height=22)
    ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:H{r}"); r += 1
    ex = ["예: 11월 작업 시작 08:00, 동절기 도장 작업 10:00 이후", "예: 동시 작업 구역(기계실)은 무전기 채널 3번으로 연락",
          "예: 가스 누출 시 바람 반대 방향 정문 집결, 대피 신호 사이렌 3회", "예: 배관 보수 수시 위험성평가 결과 공유 — 고소 작업 안전대 부착설비 추가",
          "예: 11/3~11/5 기계실 용접과 도장 작업 시간 분리(오전 용접·오후 도장)"]
    for i, it in enumerate(items):
        box(ws, f"A{r}:C{r}", f"{i + 1}. {it}", f(9.5, True, INK), LEFT, SUB); box(ws, f"D{r}:H{r}", None, f(10), LEFT_TOP)
        example(ws, r, [None, None, None, ex[i]]); ws.row_dimensions[r].height = 38; r += 1
    box(ws, f"A{r}:C{r}", "기타 안건·건의", f(9.5, True, INK), LEFT, SUB); box(ws, f"D{r}:H{r}", None, f(10), LEFT_TOP); ws.row_dimensions[r].height = 38; r += 2
    sec(r, "4. 지난 회의 결정 사항 이행 확인"); r += 1
    header_row(ws, r, ["결정 사항", "", "", "", "담당", "이행 결과", "", "상태"], height=22)
    ws.merge_cells(f"A{r}:D{r}"); ws.merge_cells(f"F{r}:G{r}"); r += 1
    p0 = r
    for k_ in range(3):
        grid(ws, r, r, 1, 8, height=22); ws.merge_cells(f"A{r}:D{r}"); ws.merge_cells(f"F{r}:G{r}")
        ws.cell(row=r, column=1).alignment = LEFT; ws.cell(row=r, column=6).alignment = LEFT; r += 1
    dv_list(ws, f"H{p0}:H{r - 1}", ["완료", "진행 중", "미이행"])
    color_text_rule(ws, f"H{p0}:H{r - 1}", "미이행", RED)
    r += 1
    sec(r, "5. 이번 회의 결정 사항과 후속 조치"); r += 1
    header_row(ws, r, ["조치 내용", "", "", "", "담당 업체", "기한", "확인자", "비고"], height=22)
    ws.merge_cells(f"A{r}:D{r}"); r += 1
    q0 = r
    for k_ in range(4):
        grid(ws, r, r, 1, 8, height=22); ws.merge_cells(f"A{r}:D{r}"); ws.cell(row=r, column=1).alignment = LEFT
        ws.cell(row=r, column=6).number_format = "yyyy-mm-dd"; r += 1
    example(ws, q0, ["예: 기계실 고소 작업 구간 안전대 부착설비 2개소 설치", None, None, None, "○○설비", "2026-11-10", "○○○"])
    legend(ws, r + 1, 8)
    ws.print_area = f"A1:H{r + 1}"

    y = sheet(wb, "연간 개최 현황", [8, 13, 12, 34, 12, 14], tab=GREEN)
    title(y, "협의체 연간 개최 현황", 6, sub="매월 1회 이상 정기 회의 — 지난 달에 회의일이 비어 있으면 '미개최 확인'으로 표시됩니다")
    box(y, "A4", "연도", f(9.5, True, NAVY), CENTER, HEAD); box(y, "B4", None, f(11, True), CENTER); y.row_dimensions[4].height = 22
    box(y, "C4:F4", "← 연도를 넣으면(예: 2026) 지난 달 중 회의일이 빈 달이 '미개최 확인'으로 표시됩니다", f(8.5, False, MUTED), LEFT)
    dv_whole(y, "B4", 2020, 2100, prompt="연도 4자리")
    header_row(y, 6, ["월", "회의일", "참석 업체 수", "주요 협의·결정", "회의록 보관", "상태(자동)"], height=24)
    for i in range(12):
        x = 7 + i
        grid(y, x, x, 1, 6, height=24)
        y.cell(row=x, column=1, value=i + 1).number_format = '0"월"'
        y.cell(row=x, column=2).number_format = "yyyy-mm-dd"; y.cell(row=x, column=4).alignment = LEFT
        y.cell(row=x, column=6, value=f'=IF(B{x}<>"","개최",IF($B$4="","",IF(EOMONTH(DATE($B$4,A{x},1),0)<TODAY(),"미개최 확인","예정")))').fill = fill(SUB)
    dv_list(y, "E7:E18", ["보관", "미보관"])
    color_text_rule(y, "F7:F18", "미개최", RED)
    box(y, "A20:C20", "개최 횟수", f(9.5, True, NAVY), CENTER, HEAD); box(y, "D20", '=COUNTIF(F7:F18,"개최")', f(11, True), CENTER)
    box(y, "A21:C21", "미개최 확인 필요", f(9.5, True, NAVY), CENTER, HEAD); box(y, "D21", '=COUNTIF(F7:F18,"미개최 확인")', f(11, True, RED), CENTER)
    y.print_area = "A1:F21"

    guide_sheet(wb, "안전 및 보건에 관한 협의체 회의록", [
        ("누가·언제", ["관계수급인 근로자가 도급인의 사업장에서 작업하는 경우 도급인은 도급인과 수급인을 구성원으로 하는 협의체를 구성·운영해야 합니다(법 제64조제1항제1호). 협의체는 도급인 및 그의 수급인 전원으로 구성합니다(시행규칙 제79조제1항).",
                    "매월 1회 이상 정기적으로 회의를 열고 그 결과를 기록·보존해야 합니다(제79조제3항). '연간 개최 현황' 시트로 빠진 달이 없는지 확인하세요."]),
        ("무엇을 협의하나", ["제79조제2항의 다섯 가지(작업 시작 시간, 작업·작업장 간 연락방법, 재해 위험 시 대피방법, 위험성평가 실시, 연락 방법 및 작업공정 조정)를 매 회의에서 다루고 결과를 적습니다. 회색 칸은 원문입니다.",
                         "결정 사항은 담당 업체와 기한을 정해 다음 회의에서 이행 결과를 확인하세요."]),
        ("함께 쓰면 좋은 것", ["도급인의 작업장 순회점검(시행규칙 제80조: 건설업·제조업 등 2일에 1회 이상, 그 밖의 사업 1주일에 1회 이상)과 합동 안전·보건점검(제82조) 결과를 협의 안건으로 올리면 연결이 됩니다."]),
    ], source="근거: 산업안전보건법 제64조, 시행규칙 제79조·제80조·제82조. 화면으로 작성: safeplum.com/tools/council/")
    wb.move_sheet("작성 안내", offset=-2); wb.active = 1
    p = OUT / "SafePlum_council_minutes.xlsx"; wb.save(p); return p


if __name__ == "__main__":
    for fn in (risk, tbm, patrol, edu, near_miss, ppe, workplan, confined, hot_work, council):
        print(fn())
